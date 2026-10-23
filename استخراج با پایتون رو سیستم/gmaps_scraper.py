#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
================================================================================
🗺️ استخراج‌کنندهٔ حرفه‌ای و تک‌فایلی اطلاعات کسب‌وکارها از گوگل مپ (نسخه ۳.۰)
Single-File Standalone Google Maps Business Scraper (Web Panel + CLI)
================================================================================

تمامی اجزای پروژه (موتور Playwright، سرور وب پایدار Waitress/Flask، پایگاه‌دادهٔ
تک‌فایلی SQLite برای ذخیرهٔ پیشرفت و ادامهٔ کار، خروجی درلحظهٔ اکسل/CSV در حافظه،
و کل رابط کاربری وب فارسی RTL) به‌صورت کاملاً مستقل در همین یک فایل پایتون
یکپارچه شده‌اند و به هیچ فایل کد/قالب اضافه‌ای نیاز نیست.

نحوهٔ اجرا:
    python gmaps_scraper.py

گزینه‌های خط فرمان (اختیاری):
    python gmaps_scraper.py --port 5000          # تغییر پورت وب‌سرور
    python gmaps_scraper.py --no-browser         # عدم باز کردن خودکار مرورگر در سیستم محلی
    python gmaps_scraper.py --install-browser    # نصب خودکار مرورگر Chromium و خروج

متغیرهای محیطی پشتیبانی‌شده (مخصوص داکر / Railway / سرور لینوکس):
    PORT=5000            پورت اجرای وب‌سرور (پیش‌فرض: 5000)
    DATA_DIR=/app/data   مسیر ذخیرهٔ فایل وضعیت/دیتابیس (برای Volume در Railway)
    NO_AUTO_INSTALL=1    غیرفعال‌سازی نصب خودکار پکیج‌های کسری با pip
================================================================================
"""

from __future__ import annotations

import argparse
import csv
import gc
import io
import json
import os
import platform
import random
import re
import sqlite3
import subprocess
import sys
import threading
import time
import traceback
import urllib.parse
import webbrowser
from pathlib import Path

# ── بهینه‌سازی ۱: کاهش حجم استک Threadها از ۸MB به ۵۱۲KB ─────────────────────
# حتماً پیش از ساخت هرگونه Thread اجرا می‌شود تا در اجراهای چندروزه روی سرورهای
# ابری (مانند Railway/Docker) هرگز خطای «RuntimeError: can't start new thread» رخ ندهد.
try:
    threading.stack_size(512 * 1024)
except (ValueError, RuntimeError):
    pass


# ── بهینه‌سازی ۲: بررسی و نصب خودکار وابستگی‌های پایتون در صورت نیاز ─────────
REQUIRED_PACKAGES = [
    ("flask", "flask>=3.0,<4.0"),
    ("openpyxl", "openpyxl>=3.1,<4.0"),
    ("playwright", "playwright>=1.45.0"),
    ("waitress", "waitress>=3.0,<4.0"),
]


def ensure_python_dependencies() -> None:
    """بررسی می‌کند آیا کتابخانه‌های موردنیاز نصب هستند یا نه؛ در صورت کسری و
    مجاز بودن، آن‌ها را خودکار با pip نصب می‌کند تا نیاز به requirements.txt نباشد."""
    missing = []
    for mod_name, pip_spec in REQUIRED_PACKAGES:
        try:
            __import__(mod_name)
        except ImportError:
            missing.append(pip_spec)

    if not missing:
        return

    if os.environ.get("NO_AUTO_INSTALL") == "1":
        sys.stderr.write(
            "پکیج‌های زیر نصب نیستند:\n  "
            + "\n  ".join(missing)
            + f"\nلطفاً دستور زیر را اجرا کنید:\n  {sys.executable} -m pip install "
            + " ".join(f'"{m}"' for m in missing)
            + "\n"
        )
        sys.exit(1)

    print(
        f"[راه‌اندازی اولیه] در حال نصب خودکار پکیج‌های موردنیاز ({', '.join(missing)})...",
        flush=True,
    )
    try:
        subprocess.check_call(
            [sys.executable, "-m", "pip", "install", *missing]
        )
    except Exception as exc:
        sys.stderr.write(
            f"\n[خطا] نصب خودکار پکیج‌ها ناموفق بود: {exc}\n"
            f"لطفاً دستی اجرا کنید:\n  {sys.executable} -m pip install flask openpyxl playwright waitress\n"
        )
        sys.exit(1)


ensure_python_dependencies()

from flask import Flask, Response, jsonify, request, send_file  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from playwright.sync_api import sync_playwright  # noqa: E402

# ==============================================================================
# ثوابت و مسیر ذخیره‌سازی
# ==============================================================================

APP_VERSION = "3.0 (نسخهٔ تک‌فایلی یکپارچه و بهینه‌شده)"
PROCESS_START = time.time()
MAX_KEYWORDS = 30

BASE_DIR = Path(__file__).resolve().parent

# تعیین هوشمند مسیر فایل دیتابیس:
# ۱. اگر متغیر محیطی DATA_DIR داده شده باشد، از همان استفاده می‌شود.
# ۲. اگر پوشهٔ data کنار فایل وجود داشته باشد (مثلاً Volume مانت‌شده در Railway به /app/data)،
#    دیتابیس داخل آن ذخیره می‌شود تا بین دیپلوی‌ها حفظ شود.
# ۳. در غیر این صورت، بدون ساخت هیچ پوشهٔ اضافه‌ای، مستقیم یک فایل تکی «gmaps_state.db»
#    کنار همین اسکریپت پایتون قرار می‌گیرد.
if os.environ.get("DATA_DIR"):
    STORAGE_DIR = Path(os.environ["DATA_DIR"]).resolve()
    STORAGE_DIR.mkdir(parents=True, exist_ok=True)
    DB_PATH = STORAGE_DIR / "gmaps_state.db"
elif (BASE_DIR / "data").is_dir():
    STORAGE_DIR = BASE_DIR / "data"
    DB_PATH = STORAGE_DIR / "gmaps_state.db"
else:
    STORAGE_DIR = BASE_DIR
    DB_PATH = BASE_DIR / "gmaps_state.db"


# ==============================================================================
# توابع کمکی پردازش متن، اعداد فارسی، تلفن و لوکیشن‌ها
# ==============================================================================

_PERSIAN_ARABIC_DIGITS = str.maketrans(
    "۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩٫",
    "01234567890123456789.",
)

# حذف کاراکترهای نامرئی جهت‌دهی یونیکد و آیکون‌های Material (حفظ کامل نیم‌فاصله \u200c)
_CLEAN_TEXT_RE = re.compile(r"[\ue000-\uf8ff\u200e\u200f\u202a-\u202e\u2066-\u2069]")


def to_english_digits(text: str) -> str:
    """اعداد فارسی/عربی و ممیز فارسی را به معادل استاندارد انگلیسی تبدیل می‌کند."""
    return (text or "").translate(_PERSIAN_ARABIC_DIGITS)


def clean_display_text(text: str | None) -> str:
    """متن استخراج‌شده از گوگل مپ را از کاراکترهای کنترلی BiDi و آیکون‌ها پاک می‌کند
    تا در اکسل هیچ کاراکتر مربعی (□) یا به‌هم‌ریختگی جهت متن ایجاد نشود."""
    if not text:
        return ""
    cleaned = _CLEAN_TEXT_RE.sub("", text)
    cleaned = re.sub(r"[ \t\r\n]+", " ", cleaned).strip()
    return cleaned


def normalize_phone(raw: str) -> tuple[str | None, str]:
    """رشتهٔ تلفن خام (با ارقام فارسی، عربی یا انگلیسی و پیش‌شماره +98 / 0098) را
    نرمال‌سازی کرده و نوع آن را به‌صورت ('mobile' | 'landline' | None, digits) برمی‌گرداند."""
    ascii_raw = to_english_digits(raw or "")
    digits = re.sub(r"\D", "", ascii_raw)
    if not digits:
        return None, ""

    if digits.startswith("0098") and len(digits) > 6:
        digits = "0" + digits[4:]
    elif digits.startswith("98") and len(digits) > 10:
        digits = "0" + digits[2:]
    elif len(digits) == 10 and digits.startswith("9"):
        digits = "0" + digits

    if len(digits) < 4:
        return None, ""

    if digits.startswith("09") and len(digits) == 11:
        return "mobile", digits
    return "landline", digits


def extract_place_id_from_url(url: str | None) -> str | None:
    """شناسهٔ یکتا و همیشگی مکان در گوگل مپ (به فرم 0x...:0x...) را مستقیماً از
    آدرس استخراج می‌کند. این شناسه حتی در لینک‌های خام لیست نتایج قبل از باز کردن
    صفحه وجود دارد و امکان رد کردن فوق‌سریع تکراری‌ها را فراهم می‌کند."""
    if not url:
        return None
    m = re.search(r"!1s(0x[0-9a-fA-F]+:0x[0-9a-fA-F]+)!", url)
    if m:
        return m.group(1).lower()
    m_chij = re.search(r"!19s(ChIJ[A-Za-z0-9_-]+)", url)
    if m_chij:
        return m_chij.group(1)
    return None


def extract_coords_from_url(url: str | None) -> tuple[str, str]:
    """مختصات دقیق پین کسب‌وکار (!3dLAT!4dLNG) را بدون کلیک اضافی از URL می‌خواند."""
    if not url:
        return "", ""
    m = re.search(r"!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)", url)
    if m:
        return m.group(1), m.group(2)
    m_at = re.search(r"@(-?\d+\.\d+),(-?\d+\.\d+)", url)
    if m_at:
        return m_at.group(1), m_at.group(2)
    return "", ""


def build_search_keyword(keywords: list[str]) -> str:
    """از لیست کلمات کلیدی، عبارت استاندارد جستجوی گوگل مپ می‌سازد.
    تک‌کلمه‌ای بدون پرانتز و چندکلمه‌ای به‌صورت (کلمه۱ OR کلمه۲ OR ...)."""
    cleaned = [k.strip() for k in keywords if k and k.strip()]
    if not cleaned:
        return ""
    if len(cleaned) == 1:
        return cleaned[0]
    return "(" + " OR ".join(cleaned) + ")"


def optimize_locations(locations: list[str]) -> list[str]:
    """لیست لوکیشن‌ها را بهینه‌سازی می‌کند:
    - جایگزینی کاراکتر «/» با فاصلهٔ ساده برای درک طبیعی‌تر توسط گوگل مپ
    - حذف فاصله‌های تکراری و خطوط خالی
    - حذف ردیف‌های کاملاً تکراری با حفظ ترتیب اولیه"""
    seen = set()
    result = []
    for loc in locations:
        if not loc:
            continue
        cleaned = str(loc).replace("/", " ")
        cleaned = re.sub(r"\s+", " ", cleaned).strip()
        if not cleaned:
            continue
        key = cleaned.lower()
        if key in seen:
            continue
        seen.add(key)
        result.append(cleaned)
    return result


# ==============================================================================
# موتور ذخیره‌سازی تک‌فایلی SQLite (مقاوم در برابر قطعی برق/کرش + شمارش O(1))
# ==============================================================================

DEFAULT_SETTINGS = {
    "keywords": [],
    "locations": [],
    "location_file_basename": "",
    "max_results": 150,
    "scroll_pause": 1.3,
    "item_delay": 1.2,
    "between_location_delay": 4.0,
    "skip_existing": True,
    "block_media": True,
    "headless": True,
}


class StorageManager:
    """مدیریت یکپارچهٔ تنظیمات، وضعیت اجرای کار (Resume) و نتایج استخراج‌شده
    در یک فایل واحد SQLite همراه با ایندکس درون‌حافظه‌ای O(1)."""

    def __init__(self, db_path: Path):
        self.db_path = db_path
        self._lock = threading.RLock()
        self._seen_keys: set[str] = set()
        self._total_count = 0
        self._mobile_count = 0
        self._phone_count = 0
        self._init_db()
        self._migrate_legacy_files_if_needed()
        self._load_memory_index()

    def _connect(self) -> sqlite3.Connection:
        conn = sqlite3.connect(str(self.db_path), timeout=30.0)
        conn.row_factory = sqlite3.Row
        conn.execute("PRAGMA journal_mode=DELETE;")
        conn.execute("PRAGMA synchronous=NORMAL;")
        return conn

    def _init_db(self) -> None:
        with self._lock, self._connect() as conn:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS kv_store (
                    key   TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );

                CREATE TABLE IF NOT EXISTS results (
                    id          INTEGER PRIMARY KEY AUTOINCREMENT,
                    dedup_key   TEXT NOT NULL UNIQUE,
                    place_id    TEXT,
                    location    TEXT,
                    query       TEXT,
                    name        TEXT,
                    category    TEXT,
                    address     TEXT,
                    mobiles     TEXT,
                    landlines   TEXT,
                    website     TEXT,
                    rating      TEXT,
                    lat         TEXT,
                    lng         TEXT,
                    url         TEXT,
                    created_at  TEXT
                );
                """
            )
            conn.commit()

    def _migrate_legacy_files_if_needed(self) -> None:
        """اگر فایل‌های نسخه‌های قدیمی (settings.json / job_state.json / results.jsonl)
        در پوشهٔ داده موجود باشند و دیتابیس تازه ساخته شده باشد، داده‌ها را خودکار
        منتقل می‌کند تا هیچ اطلاعات قبلی روی سرور/سیستم کاربر از دست نرود."""
        legacy_dir = self.db_path.parent
        legacy_settings = legacy_dir / "settings.json"
        legacy_state = legacy_dir / "job_state.json"
        legacy_results = legacy_dir / "results.jsonl"

        with self._lock, self._connect() as conn:
            has_kv = conn.execute("SELECT COUNT(*) FROM kv_store").fetchone()[0] > 0
            has_res = conn.execute("SELECT COUNT(*) FROM results").fetchone()[0] > 0

            if not has_kv and legacy_settings.exists():
                try:
                    data = json.loads(legacy_settings.read_text(encoding="utf-8"))
                    merged = dict(DEFAULT_SETTINGS)
                    merged["keywords"] = [str(k).strip() for k in data.get("keywords", []) if str(k).strip()]
                    merged["locations"] = [str(l).strip() for l in data.get("locations", []) if str(l).strip()]
                    merged["location_file_basename"] = str(data.get("location_file_basename", "") or "").strip()
                    conn.execute(
                        "INSERT OR REPLACE INTO kv_store (key, value) VALUES (?, ?)",
                        ("settings", json.dumps(merged, ensure_ascii=False)),
                    )
                except Exception:
                    pass

            if not has_kv and legacy_state.exists():
                try:
                    st = json.loads(legacy_state.read_text(encoding="utf-8"))
                    conn.execute(
                        "INSERT OR REPLACE INTO kv_store (key, value) VALUES (?, ?)",
                        ("job_state", json.dumps(st, ensure_ascii=False)),
                    )
                except Exception:
                    pass

            if not has_res and legacy_results.exists():
                try:
                    with open(legacy_results, "r", encoding="utf-8") as f:
                        for line in f:
                            line = line.strip()
                            if not line:
                                continue
                            try:
                                rec = json.loads(line)
                            except Exception:
                                continue
                            key = (rec.get("place_id") or rec.get("url") or "").strip()
                            if not key:
                                continue
                            lat, lng = extract_coords_from_url(rec.get("url"))
                            conn.execute(
                                """
                                INSERT OR IGNORE INTO results
                                (dedup_key, place_id, location, query, name, category, address,
                                 mobiles, landlines, website, rating, lat, lng, url, created_at)
                                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                                """,
                                (
                                    key,
                                    rec.get("place_id") or "",
                                    rec.get("location") or "",
                                    rec.get("query") or "",
                                    rec.get("name") or "",
                                    rec.get("category") or "",
                                    rec.get("address") or "",
                                    json.dumps(rec.get("mobiles") or [], ensure_ascii=False),
                                    json.dumps(rec.get("landlines") or [], ensure_ascii=False),
                                    rec.get("website") or "",
                                    rec.get("rating") or "",
                                    lat,
                                    lng,
                                    rec.get("url") or "",
                                    time.strftime("%Y-%m-%d %H:%M:%S"),
                                ),
                            )
                except Exception:
                    pass
            conn.commit()

    def _load_memory_index(self) -> None:
        with self._lock, self._connect() as conn:
            self._seen_keys.clear()
            self._total_count = 0
            self._mobile_count = 0
            self._phone_count = 0
            rows = conn.execute(
                "SELECT dedup_key, mobiles, landlines FROM results"
            ).fetchall()
            for r in rows:
                self._seen_keys.add(r["dedup_key"])
                self._total_count += 1
                mobs = json.loads(r["mobiles"] or "[]")
                lands = json.loads(r["landlines"] or "[]")
                if mobs:
                    self._mobile_count += 1
                if mobs or lands:
                    self._phone_count += 1

    # ── شمارنده‌ها و بررسی یکتا بودن در حافظه O(1) ──

    @property
    def total_count(self) -> int:
        return self._total_count

    @property
    def mobile_count(self) -> int:
        return self._mobile_count

    @property
    def phone_count(self) -> int:
        return self._phone_count

    def has_seen(self, key: str | None) -> bool:
        if not key:
            return False
        return key in self._seen_keys

    # ── مدیریت تنظیمات ──

    def get_settings(self) -> dict:
        with self._lock, self._connect() as conn:
            row = conn.execute(
                "SELECT value FROM kv_store WHERE key = 'settings'"
            ).fetchone()
        result = dict(DEFAULT_SETTINGS)
        if row:
            try:
                data = json.loads(row["value"])
                result["keywords"] = [str(k).strip() for k in data.get("keywords", []) if str(k).strip()]
                result["locations"] = [str(l).strip() for l in data.get("locations", []) if str(l).strip()]
                result["location_file_basename"] = str(data.get("location_file_basename", "") or "").strip()
                for num_key, cast in (
                    ("max_results", int),
                    ("scroll_pause", float),
                    ("item_delay", float),
                    ("between_location_delay", float),
                ):
                    if data.get(num_key) is not None:
                        result[num_key] = cast(data[num_key])
                for bool_key in ("skip_existing", "block_media", "headless"):
                    if bool_key in data:
                        result[bool_key] = bool(data[bool_key])
            except Exception:
                pass
        return result

    def save_settings(self, settings: dict) -> None:
        with self._lock, self._connect() as conn:
            conn.execute(
                "INSERT OR REPLACE INTO kv_store (key, value) VALUES (?, ?)",
                ("settings", json.dumps(settings, ensure_ascii=False)),
            )
            conn.commit()

    # ── مدیریت وضعیت کار (Resume) ──

    def get_job_state(self) -> dict | None:
        with self._lock, self._connect() as conn:
            row = conn.execute(
                "SELECT value FROM kv_store WHERE key = 'job_state'"
            ).fetchone()
        if not row:
            return None
        try:
            return json.loads(row["value"])
        except Exception:
            return None

    def save_job_state(self, state_dict: dict) -> None:
        snapshot = dict(state_dict)
        snapshot["logs"] = []
        with self._lock, self._connect() as conn:
            conn.execute(
                "INSERT OR REPLACE INTO kv_store (key, value) VALUES (?, ?)",
                ("job_state", json.dumps(snapshot, ensure_ascii=False)),
            )
            conn.commit()

    # ── افزودن، خواندن و پاک‌سازی نتایج ──

    def append_result(self, record: dict) -> bool:
        """یک رکورد کسب‌وکار را ذخیره می‌کند. اگر تکراری نباشد True برمی‌گرداند."""
        key = (record.get("place_id") or record.get("url") or "").strip()
        if not key:
            return False

        with self._lock:
            if key in self._seen_keys:
                return False

            mobiles = record.get("mobiles") or []
            landlines = record.get("landlines") or []
            now_str = time.strftime("%Y-%m-%d %H:%M:%S")

            with self._connect() as conn:
                cur = conn.execute(
                    """
                    INSERT OR IGNORE INTO results
                    (dedup_key, place_id, location, query, name, category, address,
                     mobiles, landlines, website, rating, lat, lng, url, created_at)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        key,
                        record.get("place_id") or "",
                        record.get("location") or "",
                        record.get("query") or "",
                        record.get("name") or "",
                        record.get("category") or "",
                        record.get("address") or "",
                        json.dumps(mobiles, ensure_ascii=False),
                        json.dumps(landlines, ensure_ascii=False),
                        record.get("website") or "",
                        record.get("rating") or "",
                        record.get("lat") or "",
                        record.get("lng") or "",
                        record.get("url") or "",
                        now_str,
                    ),
                )
                conn.commit()
                if cur.rowcount == 0:
                    self._seen_keys.add(key)
                    return False

            self._seen_keys.add(key)
            self._total_count += 1
            if mobiles:
                self._mobile_count += 1
            if mobiles or landlines:
                self._phone_count += 1
            return True

    def _row_to_dict(self, r: sqlite3.Row) -> dict:
        return {
            "id": r["id"],
            "place_id": r["place_id"],
            "location": r["location"],
            "query": r["query"],
            "name": r["name"],
            "category": r["category"],
            "address": r["address"],
            "mobiles": json.loads(r["mobiles"] or "[]"),
            "landlines": json.loads(r["landlines"] or "[]"),
            "website": r["website"],
            "rating": r["rating"],
            "lat": r["lat"],
            "lng": r["lng"],
            "url": r["url"],
            "created_at": r["created_at"],
        }

    def get_recent_results(self, limit: int = 30, search_q: str = "") -> list[dict]:
        q = (search_q or "").strip()
        with self._lock, self._connect() as conn:
            if q:
                like = f"%{q}%"
                rows = conn.execute(
                    """
                    SELECT * FROM results
                    WHERE name LIKE ? OR location LIKE ? OR address LIKE ?
                       OR category LIKE ? OR mobiles LIKE ? OR landlines LIKE ?
                    ORDER BY id DESC LIMIT ?
                    """,
                    (like, like, like, like, like, like, limit),
                ).fetchall()
            else:
                rows = conn.execute(
                    "SELECT * FROM results ORDER BY id DESC LIMIT ?",
                    (limit,),
                ).fetchall()
        return [self._row_to_dict(r) for r in rows]

    def load_all_results(self, filter_mode: str = "all") -> list[dict]:
        """بارگذاری نتایج برای خروجی اکسل/CSV با پشتیبانی از فیلتر شماره تماس."""
        with self._lock, self._connect() as conn:
            rows = conn.execute("SELECT * FROM results ORDER BY id ASC").fetchall()
        records = [self._row_to_dict(r) for r in rows]
        if filter_mode == "with_mobile":
            return [r for r in records if r["mobiles"]]
        if filter_mode == "with_phone":
            return [r for r in records if r["mobiles"] or r["landlines"]]
        return records

    def clear_results_only(self) -> None:
        with self._lock, self._connect() as conn:
            conn.execute("DELETE FROM results")
            conn.commit()
            conn.execute("VACUUM")
            self._seen_keys.clear()
            self._total_count = 0
            self._mobile_count = 0
            self._phone_count = 0

    def reset_everything(self) -> None:
        with self._lock, self._connect() as conn:
            conn.execute("DELETE FROM results")
            conn.execute("DELETE FROM kv_store")
            conn.commit()
            conn.execute("VACUUM")
            self._seen_keys.clear()
            self._total_count = 0
            self._mobile_count = 0
            self._phone_count = 0


# ==============================================================================
# مانیتورینگ منابع سیستم (RAM / Thread / Chromium / Cgroup)
# ==============================================================================

def _read_int_file(path: str) -> int | None:
    try:
        with open(path, "r", encoding="utf-8") as f:
            return int(f.read().strip())
    except Exception:
        return None


def health_snapshot(storage: StorageManager | None = None) -> dict:
    """وضعیت لحظه‌ای منابع پروسه برای پایش پایداری روی سرور یا کامپیوتر شخصی."""
    rss_kb = 0
    try:
        with open("/proc/self/status", "r", encoding="utf-8") as f:
            for line in f:
                if line.startswith("VmRSS"):
                    rss_kb = int(line.split()[1])
                    break
    except Exception:
        pass

    mem_used = _read_int_file("/sys/fs/cgroup/memory.current")
    mem_limit = _read_int_file("/sys/fs/cgroup/memory.max")
    if mem_used is None:
        mem_used = _read_int_file("/sys/fs/cgroup/memory/memory.usage_in_bytes")
        mem_limit = _read_int_file("/sys/fs/cgroup/memory/memory.limit_in_bytes")
    if mem_limit is not None and mem_limit >= 2**60:
        mem_limit = None

    chrome = node = 0
    try:
        for entry in os.listdir("/proc"):
            if not entry.isdigit():
                continue
            try:
                with open(f"/proc/{entry}/cmdline", "rb") as f:
                    cmd = f.read()
            except Exception:
                continue
            if b"chrome" in cmd or b"headless_shell" in cmd:
                chrome += 1
            elif cmd.startswith(b"node") or b"/node" in cmd[:200]:
                node += 1
    except Exception:
        pass

    return {
        "threads": threading.active_count(),
        "rss_mb": round(rss_kb / 1024, 1) if rss_kb else None,
        "cgroup_mem_used_mb": round(mem_used / 1048576, 1) if mem_used else None,
        "cgroup_mem_limit_mb": round(mem_limit / 1048576, 1) if mem_limit else None,
        "chrome_procs": chrome,
        "node_procs": node,
        "results_count": storage.total_count if storage else 0,
        "mobile_count": storage.mobile_count if storage else 0,
        "phone_count": storage.phone_count if storage else 0,
        "uptime_hours": round((time.time() - PROCESS_START) / 3600, 2),
    }


def health_log_line(storage: StorageManager | None = None, prefix: str = "health") -> str:
    snap = health_snapshot(storage)
    rss_part = f"{snap['rss_mb']}MB" if snap["rss_mb"] is not None else "N/A"
    mem_part = (
        f"{snap['cgroup_mem_used_mb']}/{snap['cgroup_mem_limit_mb']}MB"
        if snap["cgroup_mem_used_mb"] is not None
        else "host"
    )
    return (
        f"[{prefix}] threads={snap['threads']} rss={rss_part} "
        f"mem={mem_part} chrome={snap['chrome_procs']} node={snap['node_procs']} "
        f"results={snap['results_count']} mobiles={snap['mobile_count']} "
        f"uptime={snap['uptime_hours']}h"
    )


# ==============================================================================
# موتور اصلی استخراج از گوگل مپ (Playwright — چرخهٔ سه‌لایه و ضدکرش)
# ==============================================================================

# عبارت‌های دقیق تشخیص صفحهٔ کپچا/مسدودیت گوگل (بدون خطای مثبت کاذب روی نام فروشگاه‌ها)
BLOCK_PHRASES = (
    "unusual traffic from your computer network",
    "our systems have detected unusual traffic",
    "به نظر می‌رسد که فعالیت غیرمعمولی",
    "سیستم‌های ما ترافیک غیرعادی",
)

# پشتیبانی همزمان از رابط فارسی و انگلیسی گوگل مپ برای تشخیص پایان فهرست و عدم وجود نتیجه
END_OF_LIST_MARKERS = (
    "به‌پایان فهرست رسیده‌اید",
    "به پایان فهرست رسیده‌اید",
    "you've reached the end of the list",
    "reached the end of the list",
)

NOT_FOUND_MARKERS = (
    "گوگل مپ نمی‌تواند",
    "را پیدا کند",
    "چیزی پیدا نشد",
    "google maps can't find",
    "no results found",
)

CRASH_MARKERS = (
    "crashed",
    "target closed",
    "has been closed",
    "context was destroyed",
    "browser has been closed",
    "connection closed",
)

# کوکی استاندارد پذیرش رضایت‌نامهٔ گوگل (جلوگیری از قفل شدن روی consent.google.com در سرورهای اروپا)
GOOGLE_CONSENT_COOKIES = [
    {
        "name": "SOCS",
        "value": "CAISHAgBEhJnd3NfMjAyNDAxMDEtMF9SQzIaAmZhIAEaBgiA_LyuBg",
        "domain": ".google.com",
        "path": "/",
    }
]


def _is_crash_error(exc: Exception) -> bool:
    msg = str(exc).lower()
    return any(marker in msg for marker in CRASH_MARKERS)


def install_playwright_chromium(log_fn=print) -> bool:
    """نصب خودکار مرورگر Chromium (و وابستگی‌های لینوکس در صورت دسترسی root)."""
    log_fn("در حال دانلود و نصب خودکار مرورگر Chromium برای Playwright (فقط بار اول)...")
    try:
        cmd = [sys.executable, "-m", "playwright", "install"]
        if sys.platform.startswith("linux") and hasattr(os, "geteuid") and os.geteuid() == 0:
            cmd.append("--with-deps")
        cmd.append("chromium")
        subprocess.check_call(cmd)
        log_fn("مرورگر Chromium با موفقیت نصب شد.")
        return True
    except Exception as exc:
        log_fn(f"نصب خودکار Chromium ناموفق بود: {exc}")
        return False


class GoogleMapsScraper:
    """موتور استخراج اطلاعات کسب‌وکارها از گوگل مپ با معماری سه‌لایه:
      لایه ۱: self._pw      -> درایور Playwright (یک بار در کل اجرای جاب)
      لایه ۲: self._browser -> مرورگر Chromium (فقط در کرش واقعی راه‌اندازی مجدد می‌شود)
      لایه ۳: self._context -> تب/کانتکست تازه هر ۲۵ آیتم برای صفر شدن مصرف RAM"""

    def __init__(
        self,
        headless: bool = True,
        block_media: bool = True,
        log_fn=None,
        seen_checker=None,
        storage_ref: StorageManager | None = None,
    ):
        self.headless = headless
        self.block_media = block_media
        self._log = log_fn or (lambda _msg: None)
        self._has_seen = seen_checker or (lambda _key: False)
        self._storage_ref = storage_ref

        self._pw = None
        self._browser = None
        self._context = None
        self._page = None
        self._stop_flag = False
        self._auto_install_tried = False

    def request_stop(self) -> None:
        self._stop_flag = True

    # ── چرخهٔ حیات سه‌لایه (درایور / مرورگر / کانتکست) ──

    def _ensure_driver(self) -> None:
        if self._pw is not None:
            return
        self._pw = sync_playwright().start()

    def _stop_driver(self) -> None:
        try:
            if self._pw:
                self._pw.stop()
        except Exception:
            pass
        self._pw = None

    def _try_launch_browser_once(self, docker_safe_args: list[str]) -> list[str]:
        errors = []
        launch_attempts = [
            {"headless": self.headless, "args": docker_safe_args},
            {"headless": self.headless, "channel": "chromium", "args": docker_safe_args},
            {"headless": self.headless, "channel": "chrome", "args": docker_safe_args},
            {"headless": self.headless, "channel": "msedge", "args": docker_safe_args},
        ]
        for attempt in launch_attempts:
            try:
                self._browser = self._pw.chromium.launch(**attempt)
                return []
            except Exception as exc:
                errors.append(f"  - {attempt.get('channel', 'پیش‌فرض')}: {exc}")
                self._browser = None
        return errors

    def _launch_browser(self) -> None:
        self._ensure_driver()
        docker_safe_args = [
            "--no-sandbox",
            "--disable-dev-shm-usage",
            "--disable-gpu",
            "--disable-extensions",
            "--disable-background-networking",
            "--disable-background-timer-throttling",
            "--disable-renderer-backgrounding",
            "--disable-blink-features=AutomationControlled",
            "--renderer-process-limit=2",
            "--js-flags=--max-old-space-size=512",
        ]

        errors = self._try_launch_browser_once(docker_safe_args)
        if self._browser is not None:
            return

        # تلاش هوشمند نصب خودکار Chromium در اولین اجرا اگر مرورگر روی سیستم نصب نبود
        if not self._auto_install_tried:
            self._auto_install_tried = True
            if install_playwright_chromium(self._log):
                errors = self._try_launch_browser_once(docker_safe_args)
                if self._browser is not None:
                    return

        self._stop_driver()
        details = "\n".join(errors)
        raise RuntimeError(
            "مرورگر Chromium پیدا یا اجرا نشد. لطفاً یکی از این راه‌ها را انجام دهید:\n"
            f"  ۱) دستور زیر را در ترمینال اجرا کنید:\n"
            f"     {sys.executable} -m playwright install --with-deps chromium\n"
            "  ۲) در ویندوز اگر آنتی‌ویروس یا تحریم‌شکن مانع دانلود شده، موقتاً بررسی کنید.\n"
            f"جزئیات تلاش‌ها:\n{details}"
        )

    def _new_context(self) -> None:
        if self._context is not None:
            try:
                self._context.close()
            except Exception:
                pass
            self._context = None
            self._page = None

        self._context = self._browser.new_context(
            locale="fa-IR",
            viewport={"width": 1366, "height": 900},
            user_agent=(
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/135.0.0.0 Safari/537.36"
            ),
            extra_http_headers={
                "Accept-Language": "fa-IR,fa;q=0.9,en-US;q=0.8,en;q=0.7"
            },
        )
        try:
            self._context.add_cookies(GOOGLE_CONSENT_COOKIES)
        except Exception:
            pass

        # مسدودسازی تصاویر/رسانه‌های سنگین نقشه برای کاهش ۶۰٪ مصرف رم و افزایش سرعت
        if self.block_media:
            self._context.route(
                "**/*",
                lambda route: (
                    route.abort()
                    if route.request.resource_type in ("image", "media", "font")
                    else route.continue_()
                ),
            )

        self._page = self._context.new_page()
        self._page.set_default_timeout(20000)

    def start(self) -> None:
        self._ensure_driver()
        self._launch_browser()
        self._new_context()

    def _restart_context(self) -> None:
        self._log("ری‌استارت سبک تب مرورگر برای آزادسازی حافظهٔ رم...")
        self._log(health_log_line(self._storage_ref, "health-context-refresh"))
        try:
            self._new_context()
        except Exception as exc:
            self._log(f"ری‌استارت سبک ناموفق بود ({exc})؛ راه‌اندازی مجدد مرورگر...")
            self._restart_browser()
            return
        gc.collect()
        time.sleep(0.4)

    def _restart_browser(self) -> None:
        self._log("در حال بستن مرورگر آسیب‌دیده و راه‌اندازی مجدد Chromium...")
        self._log(health_log_line(self._storage_ref, "health-browser-restart"))
        try:
            if self._browser:
                self._browser.close()
        except Exception:
            pass
        self._browser = None
        self._context = None
        self._page = None
        time.sleep(1.0)
        try:
            self._launch_browser()
            self._new_context()
        except Exception as exc:
            self._log(f"راه‌اندازی مجدد مرورگر ناموفق بود ({exc})؛ ری‌استارت کامل درایور...")
            self._hard_restart()
            return
        gc.collect()
        self._log("مرورگر با موفقیت مجدداً آماده شد.")

    def _hard_restart(self) -> None:
        try:
            if self._browser:
                self._browser.close()
        except Exception:
            pass
        self._stop_driver()
        self._browser = None
        self._context = None
        self._page = None
        time.sleep(2.0)
        self._ensure_driver()
        self._launch_browser()
        self._new_context()
        gc.collect()
        self._log("درایور و مرورگر به‌طور کامل از صفر راه‌اندازی شدند.")

    def stop(self) -> None:
        try:
            if self._context:
                self._context.close()
        except Exception:
            pass
        try:
            if self._browser:
                self._browser.close()
        except Exception:
            pass
        self._stop_driver()
        self._browser = None
        self._context = None
        self._page = None

    # ── تشخیص صفحهٔ رضایت‌نامهٔ کوکی گوگل، کپچا و استخراج فیلدها ──

    def _dismiss_consent_if_needed(self) -> None:
        """در صورتی که سرور در اروپا باشد و گوگل به consent.google.com هدایت کند،
        فرم رضایت کوکی را خودکار تأیید می‌کند."""
        try:
            cur_url = (self._page.url or "").lower()
            if "consent.google." not in cur_url:
                return
            self._log("صفحهٔ تأیید کوکی گوگل تشخیص داده شد؛ در حال عبور خودکار...")
            selectors = [
                'form[action*="save"] button',
                'button[aria-label*="Accept"]',
                'button[aria-label*="پذیرش"]',
                'button:has-text("Accept all")',
                'button:has-text("پذیرش همه")',
                'button:has-text("Alles accepteren")',
                'button:has-text("Alle akzeptieren")',
            ]
            for sel in selectors:
                btn = self._page.query_selector(sel)
                if btn:
                    btn.click()
                    self._page.wait_for_load_state("domcontentloaded", timeout=15000)
                    time.sleep(1.5)
                    break
        except Exception:
            pass

    def _body_text(self) -> str:
        try:
            return self._page.inner_text("body")
        except Exception:
            return ""

    def _check_blocked(self) -> bool:
        """تشخیص دقیق کپچا/مسدودیت گوگل بدون خطای کاذب روی کلماتی مثل «رباتیک»."""
        try:
            cur_url = (self._page.url or "").lower()
            if "/sorry/" in cur_url:
                return True
            if self._page.query_selector("#captcha-form, iframe[src*='recaptcha']") is not None:
                return True
            text = self._body_text().lower()
            return any(phrase in text for phrase in BLOCK_PHRASES)
        except Exception:
            return False

    def _wait_for_search_ready(self, timeout_sec: float = 9.0) -> None:
        """به‌جای صبر کورکورانه، منتظر ظاهر شدن لیست نتایج، نتیجهٔ تکی یا پیام عدم یافتن می‌ماند."""
        deadline = time.time() + timeout_sec
        while time.time() < deadline:
            if self._stop_flag:
                return
            try:
                if self._page.query_selector('div[role="feed"]') is not None:
                    return
                if self._page.query_selector("h1.DUwDvf") is not None:
                    return
                txt = self._body_text().lower()
                if any(m in txt for m in NOT_FOUND_MARKERS) or self._check_blocked():
                    return
            except Exception:
                pass
            time.sleep(0.4)

    def _wait_for_place_ready(self, timeout_sec: float = 8.0) -> None:
        deadline = time.time() + timeout_sec
        while time.time() < deadline:
            if self._stop_flag:
                return
            try:
                el = self._page.query_selector("h1.DUwDvf")
                if el and el.inner_text().strip():
                    return
                if self._check_blocked():
                    return
            except Exception:
                pass
            time.sleep(0.35)

    def _extract_current_place(self) -> dict:
        """استخراج کامل مشخصات کسب‌وکار از صفحهٔ جزئیات گوگل مپ."""
        name = ""
        try:
            name_el = self._page.query_selector("h1.DUwDvf") or self._page.query_selector("h1")
            if name_el:
                name = clean_display_text(name_el.inner_text())
        except Exception:
            pass

        category = ""
        try:
            cat_el = self._page.query_selector("button.DkEaL")
            if cat_el:
                category = clean_display_text(cat_el.inner_text())
        except Exception:
            pass

        address = ""
        try:
            addr_el = self._page.query_selector('button[data-item-id="address"]')
            if addr_el:
                aria = addr_el.get_attribute("aria-label") or ""
                raw_addr = re.sub(r"^(آدرس|Address)\s*:\s*", "", aria, flags=re.IGNORECASE)
                if not raw_addr.strip():
                    raw_addr = addr_el.inner_text() or ""
                address = clean_display_text(raw_addr)
        except Exception:
            pass

        mobiles: list[str] = []
        landlines: list[str] = []
        try:
            phone_els = self._page.query_selector_all('button[data-item-id^="phone:tel:"]')
            for pe in phone_els:
                candidates = []
                data_id = pe.get_attribute("data-item-id") or ""
                if data_id.startswith("phone:tel:"):
                    candidates.append(data_id[len("phone:tel:") :])
                aria = pe.get_attribute("aria-label") or ""
                if aria:
                    candidates.append(re.sub(r"^(تلفن|Phone)\s*:\s*", "", aria, flags=re.IGNORECASE))
                for raw in candidates:
                    kind, normalized = normalize_phone(raw)
                    if not kind or not normalized:
                        continue
                    if kind == "mobile":
                        if normalized not in mobiles:
                            mobiles.append(normalized)
                    else:
                        if normalized not in landlines:
                            landlines.append(normalized)
                    break
        except Exception:
            pass

        website = ""
        try:
            web_el = self._page.query_selector('a[data-item-id="authority"]')
            if web_el:
                website = (web_el.get_attribute("href") or "").strip()
        except Exception:
            pass

        rating = ""
        try:
            rat_el = self._page.query_selector("div.F7nice")
            if rat_el:
                raw_rat = to_english_digits(clean_display_text(rat_el.inner_text()))
                m_rat = re.search(r"(\d+(?:\.\d+)?)", raw_rat)
                if m_rat:
                    rating = m_rat.group(1)
        except Exception:
            pass

        url = self._page.url
        place_id = extract_place_id_from_url(url)
        lat, lng = extract_coords_from_url(url)

        return {
            "name": name,
            "category": category,
            "address": address,
            "mobiles": mobiles,
            "landlines": landlines,
            "website": website,
            "rating": rating,
            "lat": lat,
            "lng": lng,
            "url": url,
            "place_id": place_id,
        }

    def _goto_with_recovery(self, url: str, max_attempts: int = 2) -> None:
        last_err = None
        for attempt in range(1, max_attempts + 1):
            try:
                if self._page is None:
                    self._log("صفحهٔ مرورگر بسته شده بود؛ در حال بازیابی خودکار...")
                    try:
                        if self._browser is None:
                            self._launch_browser()
                        self._new_context()
                    except Exception:
                        self._hard_restart()
                self._page.goto(url, wait_until="domcontentloaded", timeout=30000)
                self._dismiss_consent_if_needed()
                return
            except Exception as exc:
                last_err = exc
                if (self._page is None or _is_crash_error(exc)) and attempt < max_attempts:
                    self._log(
                        f"کرش مرورگر هنگام باز کردن آدرس (تلاش {attempt})؛ راه‌اندازی مجدد و تلاش دوباره..."
                    )
                    self._restart_browser()
                    continue
                raise
        if last_err:
            raise last_err

    # ── حلقهٔ اصلی جستجو و استخراج یک لوکیشن ──

    def search_location(
        self,
        keyword: str,
        location: str,
        max_results: int = 150,
        scroll_pause: float = 1.3,
        item_delay: float = 1.2,
        skip_existing: bool = True,
        stable_rounds_to_stop: int = 4,
        on_item_extracted=None,
    ) -> int:
        """جستجوی یک لوکیشن در گوگل مپ، اسکرول کامل لیست، پیش‌پالایش آیتم‌های قبلاً
        استخراج‌شده و ثبت فوری هر کسب‌وکار در لحظهٔ استخراج (بدون انتظار تا پایان کل شهر)."""
        # ترتیب حیاتی و آزمایش‌شده: همیشه «<لوکیشن> <کلمهٔ کلیدی / گروه OR>»
        query = f"{location} {keyword}".strip()
        url = (
            "https://www.google.com/maps/search/"
            + urllib.parse.quote(query, safe="")
            + "?hl=fa"
        )

        self._log(f"در حال باز کردن جستجو: «{query}»")
        self._goto_with_recovery(url)
        self._wait_for_search_ready(timeout_sec=9.0)
        time.sleep(1.2)

        if self._check_blocked():
            raise RuntimeError(
                "گوگل مپ ترافیک غیرعادی (کپچا) تشخیص داد. چند دقیقه صبر کنید و سپس با "
                "دکمهٔ «ادامهٔ کار قبلی» ادامه دهید یا فاصلهٔ زمانی بین درخواست‌ها را بیشتر کنید."
            )

        feed = self._page.query_selector('div[role="feed"]')
        if feed is None:
            body_lower = self._body_text().lower()
            if any(marker in body_lower for marker in NOT_FOUND_MARKERS):
                self._log(f"هیچ نتیجه‌ای برای «{query}» یافت نشد.")
                return 0

            name_el = self._page.query_selector("h1.DUwDvf")
            if name_el:
                info = self._extract_current_place()
                if info["name"]:
                    info["location"] = location
                    info["query"] = query
                    saved = on_item_extracted(info) if on_item_extracted else True
                    if saved:
                        self._log(f"  [۱/۱] {info['name']} (نتیجهٔ تکی)")
                        return 1
                    self._log(f"  [۱/۱] رد شد (قبلاً ذخیره شده بود): {info['name']}")
                return 0

            self._log(f"هیچ نتیجه‌ای برای «{query}» یافت نشد.")
            return 0

        # اسکرول هوشمند فهرست تا رسیدن به انتها یا سقف max_results
        prev_count = -1
        stable_rounds = 0
        for _round in range(400):
            if self._stop_flag:
                self._log("درخواست توقف دریافت شد؛ اسکرول متوقف شد.")
                break

            links = feed.query_selector_all('a[href*="/maps/place/"]')
            count = len(links)

            if count >= max_results:
                self._log(f"به سقف {max_results} نتیجه رسیدیم؛ پایان اسکرول.")
                break

            body_lower = self._body_text().lower()
            if any(marker in body_lower for marker in END_OF_LIST_MARKERS):
                self._log(f"به انتهای فهرست گوگل مپ رسیدیم ({count} نتیجه).")
                break

            if count == prev_count:
                stable_rounds += 1
            else:
                stable_rounds = 0
            prev_count = count

            if stable_rounds >= stable_rounds_to_stop:
                self._log(f"فهرست دیگر رشد نکرد؛ پایان اسکرول ({count} نتیجه).")
                break

            try:
                if stable_rounds >= 1 and links:
                    links[-1].scroll_into_view_if_needed()
                feed.evaluate("el => { el.scrollBy(0, 2500); el.scrollTop = el.scrollHeight; }")
            except Exception:
                break
            time.sleep(scroll_pause + random.uniform(0.05, 0.45))

        links = feed.query_selector_all('a[href*="/maps/place/"]')
        total = min(len(links), max_results)
        self._log(f"«{query}»: {total} لینک کسب‌وکار جمع‌آوری شد.")

        # کپی لینک‌ها در پایتون و رها کردن DOM سنگین صفحهٔ جستجو
        place_urls = []
        for link in links[:total]:
            try:
                href = link.get_attribute("href")
                if href:
                    place_urls.append(href)
            except Exception:
                continue

        saved_in_location = 0
        skipped_precheck = 0
        seen_in_batch: set[str] = set()
        proactive_restart_every = 25
        navigations_since_restart = 0

        for idx, purl in enumerate(place_urls):
            if self._stop_flag:
                self._log("درخواست توقف دریافت شد؛ استخراج جزئیات متوقف شد.")
                break

            # پیش‌پالایش فوق‌سریع قبل از باز کردن صفحه:
            # اگر place_id همین لینک قبلاً در لوکیشن هم‌پوشان قبلی یا اجرای ناتمام قبلی
            # استخراج شده باشد، بدون باز کردن صفحه در ۰ میلی‌ثانیه از آن عبور می‌کنیم!
            pre_pid = extract_place_id_from_url(purl)
            if pre_pid:
                if pre_pid in seen_in_batch:
                    continue
                if skip_existing and self._has_seen(pre_pid):
                    skipped_precheck += 1
                    continue

            if navigations_since_restart >= proactive_restart_every:
                try:
                    self._restart_context()
                except Exception as exc:
                    self._log(f"هشدار در ری‌استارت پیشگیرانه: {exc}")
                navigations_since_restart = 0

            target_url = purl if "hl=" in purl else (purl + ("&hl=fa" if "?" in purl else "?hl=fa"))
            try:
                self._goto_with_recovery(target_url)
                navigations_since_restart += 1
                self._wait_for_place_ready(timeout_sec=7.5)
                time.sleep(max(0.3, item_delay) + random.uniform(0.0, 0.45))
            except Exception as exc:
                self._log(f"  [{idx + 1}/{total}] خطا در باز کردن صفحه: {exc}")
                continue

            try:
                if self._check_blocked():
                    raise RuntimeError(
                        "گوگل مپ در حین استخراج جزئیات کپچا نشان داد. تمام نتایج تا این لحظه "
                        "ذخیره شده‌اند؛ چند دقیقه صبر کرده و «ادامهٔ کار قبلی» را بزنید."
                    )
                info = self._extract_current_place()
            except RuntimeError:
                raise
            except Exception as exc:
                if _is_crash_error(exc):
                    self._log(f"  [{idx + 1}/{total}] کرش مرورگر حین خواندن جزئیات؛ بازیابی خودکار...")
                    try:
                        self._restart_browser()
                    except Exception as exc2:
                        self._log(f"خطا در بازیابی مرورگر: {exc2}")
                    navigations_since_restart = 0
                    continue
                self._log(f"  [{idx + 1}/{total}] خطا در خواندن اطلاعات: {exc}")
                continue

            if not info["name"]:
                self._log(f"  [{idx + 1}/{total}] رد شد: عنوان کسب‌وکار یافت نشد.")
                continue

            dedup_key = info.get("place_id") or info["url"]
            if dedup_key in seen_in_batch:
                continue
            seen_in_batch.add(dedup_key)

            info["location"] = location
            info["query"] = query

            saved = on_item_extracted(info) if on_item_extracted else True
            if saved:
                saved_in_location += 1
                phone_badge = ""
                if info["mobiles"]:
                    phone_badge = f" | 📱 {', '.join(info['mobiles'])}"
                elif info["landlines"]:
                    phone_badge = f" | ☎️ {', '.join(info['landlines'])}"
                self._log(f"  [{idx + 1}/{total}] ✅ {info['name']}{phone_badge}")
            else:
                self._log(f"  [{idx + 1}/{total}] ⏭️ تکراری (قبلاً ثبت شده): {info['name']}")

        if skipped_precheck > 0:
            self._log(
                f"⚡ صرفه‌جویی هوشمند: {skipped_precheck} مورد که قبلاً در دیتابیس ثبت شده بود "
                "بدون باز کردن مجدد صفحه رد شد."
            )

        return saved_in_location


# ==============================================================================
# تولید فایل خروجی اکسل (.xlsx) و CSV مستقیماً در حافظه (بدون فایل موقت روی دیسک)
# ==============================================================================

def build_excel_bytes(records: list[dict], layout: str = "full") -> io.BytesIO:
    """ساخت فایل اکسل راست‌به‌چپ (RTL) شکیل در حافظهٔ RAM با قابلیت انتخاب قالب:
      - full:    ۱۱ ستون کامل (شامل دسته‌بندی، موبایل، تلفن ثابت، وب‌سایت، امتیاز، مختصات)
      - classic: دقیقاً همان ۷ ستون نسخهٔ اصلی مخزن گیت‌هاب"""
    wb = Workbook()
    ws = wb.active
    ws.title = "نتایج استخراج گوگل مپ"
    ws.sheet_view.rightToLeft = True

    if layout == "classic":
        headers = [
            "ردیف",
            "لوکیشن جستجوشده",
            "نام فروشگاه/محل",
            "آدرس",
            "شماره موبایل",
            "شماره تلفن ثابت",
            "لینک گوگل مپ",
        ]
        widths = [8, 34, 30, 46, 20, 20, 44]
    else:
        headers = [
            "ردیف",
            "لوکیشن جستجوشده",
            "نام فروشگاه/محل",
            "دسته‌بندی کسب‌وکار",
            "شماره موبایل",
            "شماره تلفن ثابت",
            "آدرس",
            "وب‌سایت",
            "امتیاز",
            "مختصات (Lat, Lng)",
            "لینک گوگل مپ",
        ]
        widths = [8, 32, 30, 22, 20, 20, 46, 28, 10, 24, 42]

    ws.append(headers)
    ws.freeze_panes = "A2"

    header_fill = PatternFill(start_color="FFEA580C", end_color="FFEA580C", fill_type="solid")
    zebra_fill = PatternFill(start_color="FFFFF8F3", end_color="FFFFF8F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="FFE5E7EB"),
        right=Side(style="thin", color="FFE5E7EB"),
        top=Side(style="thin", color="FFE5E7EB"),
        bottom=Side(style="thin", color="FFE5E7EB"),
    )

    for cell in ws[1]:
        cell.font = Font(name="Tahoma", size=10, bold=True, color="FFFFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[1].height = 26

    body_font = Font(name="Tahoma", size=10)
    body_align = Alignment(horizontal="right", vertical="center", wrap_text=False)

    for i, r in enumerate(records, start=1):
        mobiles_str = "، ".join(r.get("mobiles") or [])
        landlines_str = "، ".join(r.get("landlines") or [])
        if layout == "classic":
            row_vals = [
                i,
                r.get("location", ""),
                r.get("name", ""),
                r.get("address", ""),
                mobiles_str,
                landlines_str,
                r.get("url", ""),
            ]
        else:
            coords = f"{r['lat']}, {r['lng']}" if r.get("lat") and r.get("lng") else ""
            row_vals = [
                i,
                r.get("location", ""),
                r.get("name", ""),
                r.get("category", ""),
                mobiles_str,
                landlines_str,
                r.get("address", ""),
                r.get("website", ""),
                r.get("rating", ""),
                coords,
                r.get("url", ""),
            ]
        ws.append(row_vals)
        cur_row = ws[i + 1]
        for cell in cur_row:
            cell.font = body_font
            cell.alignment = body_align
            cell.border = thin_border
            if i % 2 == 0:
                cell.fill = zebra_fill

    for col_idx, width in enumerate(widths, start=1):
        col_letter = ws.cell(row=1, column=col_idx).column_letter
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    wb.close()
    buf.seek(0)
    return buf


def build_csv_bytes(records: list[dict], layout: str = "full") -> io.BytesIO:
    """ساخت خروجی CSV با BOM استاندارد UTF-8 (باز شدن بدون به‌هم‌ریختگی حروف فارسی در اکسل)."""
    txt = io.StringIO()
    txt.write("\ufeff")
    writer = csv.writer(txt)
    if layout == "classic":
        writer.writerow([
            "ردیف", "لوکیشن جستجوشده", "نام فروشگاه/محل",
            "آدرس", "شماره موبایل", "شماره تلفن ثابت", "لینک گوگل مپ",
        ])
        for i, r in enumerate(records, start=1):
            writer.writerow([
                i,
                r.get("location", ""),
                r.get("name", ""),
                r.get("address", ""),
                "، ".join(r.get("mobiles") or []),
                "، ".join(r.get("landlines") or []),
                r.get("url", ""),
            ])
    else:
        writer.writerow([
            "ردیف", "لوکیشن جستجوشده", "نام فروشگاه/محل", "دسته‌بندی کسب‌وکار",
            "شماره موبایل", "شماره تلفن ثابت", "آدرس", "وب‌سایت",
            "امتیاز", "مختصات (Lat, Lng)", "لینک گوگل مپ",
        ])
        for i, r in enumerate(records, start=1):
            coords = f"{r['lat']}, {r['lng']}" if r.get("lat") and r.get("lng") else ""
            writer.writerow([
                i,
                r.get("location", ""),
                r.get("name", ""),
                r.get("category", ""),
                "، ".join(r.get("mobiles") or []),
                "، ".join(r.get("landlines") or []),
                r.get("address", ""),
                r.get("website", ""),
                r.get("rating", ""),
                coords,
                r.get("url", ""),
            ])

    buf = io.BytesIO(txt.getvalue().encode("utf-8"))
    buf.seek(0)
    return buf


# ==============================================================================
# قالب یکپارچهٔ وب‌پنل فارسی (HTML + CSS + JS بدون نیاز به هیچ فایل خارجی)
# ==============================================================================

HTML_TEMPLATE = r"""<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>پنل استخراج اطلاعات کسب‌وکارها از گوگل مپ</title>
<style>
    :root {
        --brand: #ea580c;
        --brand-dark: #c2410c;
        --brand-soft: #fff7ed;
        --brand-border: #fed7aa;
        --bg: #f3f4f6;
        --card: #ffffff;
        --text: #111827;
        --muted: #6b7280;
        --danger: #dc2626;
        --danger-dark: #b91c1c;
        --ok: #15803d;
        --ok-soft: #dcfce7;
    }
    * { box-sizing: border-box; }
    body {
        font-family: "Vazirmatn", "IRANSans", "Segoe UI", Tahoma, sans-serif;
        direction: rtl;
        background: var(--bg);
        margin: 0;
        color: var(--text);
        line-height: 1.6;
    }
    header {
        background: linear-gradient(135deg, #c2410c 0%, #ea580c 60%, #f97316 100%);
        color: #fff;
        padding: 24px 20px;
        text-align: center;
        box-shadow: 0 2px 10px rgba(194, 65, 12, 0.2);
    }
    header h1 { margin: 0 0 6px 0; font-size: 22px; font-weight: 800; }
    header p { margin: 0; font-size: 13px; opacity: 0.92; }
    .version-pill {
        display: inline-block;
        margin-top: 8px;
        background: rgba(255,255,255,0.18);
        padding: 2px 12px;
        border-radius: 999px;
        font-size: 11.5px;
    }

    .container { max-width: 980px; margin: 0 auto; padding: 20px 16px 44px; }

    .card {
        background: var(--card);
        border-radius: 14px;
        padding: 20px 22px;
        margin-bottom: 18px;
        box-shadow: 0 1px 4px rgba(0,0,0,0.06);
        border: 1px solid #e5e7eb;
    }
    .card h2 {
        font-size: 16.5px;
        margin: 0 0 14px 0;
        color: #7c2d12;
        border-bottom: 2px solid var(--brand-border);
        padding-bottom: 9px;
        display: flex;
        align-items: center;
        justify-content: space-between;
        gap: 10px;
        flex-wrap: wrap;
    }

    label { display: block; font-size: 13px; font-weight: 600; margin: 12px 0 5px; color: #374151; }
    input[type=text], input[type=number], select, textarea {
        width: 100%;
        padding: 9px 11px;
        border: 1px solid #d1d5db;
        border-radius: 8px;
        font-size: 13.5px;
        font-family: inherit;
        direction: rtl;
        background: #fff;
        transition: border-color 0.15s, box-shadow 0.15s;
    }
    input:focus, select:focus, textarea:focus {
        outline: none;
        border-color: var(--brand);
        box-shadow: 0 0 0 3px rgba(234, 88, 12, 0.14);
    }
    textarea { min-height: 125px; resize: vertical; }
    .hint { font-size: 11.5px; color: var(--muted); margin-top: 5px; }
    .msg-line { font-size: 12px; font-weight: 600; color: var(--brand-dark); margin-top: 6px; min-height: 18px; }

    .row { display: flex; gap: 12px; flex-wrap: wrap; }
    .row > div { flex: 1; min-width: 155px; }

    .check-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(260px, 1fr));
        gap: 10px;
        margin-top: 14px;
    }
    .check-item {
        display: flex;
        align-items: flex-start;
        gap: 8px;
        background: #f9fafb;
        border: 1px solid #e5e7eb;
        border-radius: 9px;
        padding: 9px 11px;
        font-size: 12.5px;
        cursor: pointer;
    }
    .check-item input[type=checkbox] {
        margin-top: 3px;
        accent-color: var(--brand);
        width: 16px;
        height: 16px;
        cursor: pointer;
    }

    button {
        cursor: pointer;
        border: none;
        border-radius: 8px;
        padding: 10px 18px;
        font-size: 13.5px;
        font-weight: 700;
        font-family: inherit;
        transition: background 0.15s, transform 0.05s, opacity 0.15s;
    }
    button:active { transform: translateY(1px); }
    button:disabled { opacity: 0.55; cursor: not-allowed; }
    .btn-primary { background: var(--brand); color: #fff; }
    .btn-primary:hover:not(:disabled) { background: var(--brand-dark); }
    .btn-success { background: #15803d; color: #fff; }
    .btn-success:hover:not(:disabled) { background: #166534; }
    .btn-danger { background: var(--danger); color: #fff; }
    .btn-danger:hover:not(:disabled) { background: var(--danger-dark); }
    .btn-secondary { background: #e5e7eb; color: #1f2937; }
    .btn-secondary:hover:not(:disabled) { background: #d1d5db; }
    .btn-sm { padding: 7px 13px; font-size: 12.5px; }
    .btn-row { display: flex; gap: 9px; flex-wrap: wrap; margin-top: 14px; align-items: center; }

    .status-header {
        display: flex;
        justify-content: space-between;
        align-items: center;
        flex-wrap: wrap;
        gap: 8px;
        margin-bottom: 12px;
    }
    .health-pill {
        font-size: 11.5px;
        color: #4b5563;
        background: #f3f4f6;
        padding: 4px 10px;
        border-radius: 999px;
        direction: ltr;
    }

    .status-box {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(145px, 1fr));
        gap: 11px;
        margin-bottom: 14px;
    }
    .stat {
        background: var(--brand-soft);
        border: 1px solid var(--brand-border);
        border-radius: 10px;
        padding: 11px 12px;
        text-align: center;
    }
    .stat .num { font-size: 21px; font-weight: 800; color: var(--brand-dark); }
    .stat .lbl { font-size: 11.5px; color: #6b7280; margin-top: 2px; }

    .current-loc-banner {
        background: #f8fafc;
        border: 1px dashed #cbd5e1;
        border-radius: 8px;
        padding: 8px 12px;
        font-size: 12.5px;
        color: #334155;
        margin-bottom: 10px;
    }

    .progress-bar-outer {
        background: #e5e7eb;
        border-radius: 999px;
        height: 16px;
        overflow: hidden;
        margin-bottom: 14px;
    }
    .progress-bar-inner {
        background: linear-gradient(90deg, #ea580c, #fb923c);
        height: 100%;
        width: 0%;
        transition: width 0.35s ease;
    }

    #logBox {
        background: #0f172a;
        color: #e2e8f0;
        border-radius: 10px;
        padding: 12px 14px;
        font-family: "Consolas", "Courier New", monospace;
        font-size: 12px;
        height: 270px;
        overflow-y: auto;
        white-space: pre-wrap;
        direction: rtl;
        text-align: right;
        line-height: 1.7;
    }

    .badge {
        display: inline-block;
        padding: 3px 12px;
        border-radius: 999px;
        font-size: 12px;
        font-weight: 700;
    }
    .badge.running { background: var(--ok-soft); color: var(--ok); }
    .badge.stopping { background: #fef3c7; color: #b45309; }
    .badge.idle { background: #e5e7eb; color: #4b5563; }
    .badge.error { background: #fee2e2; color: #b91c1c; }

    .warn-box {
        background: #fffbeb;
        border-right: 4px solid #f59e0b;
        border-radius: 8px;
        padding: 10px 14px;
        font-size: 12.5px;
        margin-bottom: 14px;
        color: #78350f;
    }
    .error-box {
        background: #fef2f2;
        border-right: 4px solid var(--danger);
        color: #991b1b;
    }

    .upload-box {
        background: #f9fafb;
        border: 1px dashed #cbd5e1;
        border-radius: 10px;
        padding: 12px 14px;
        margin-top: 14px;
    }

    .table-wrap {
        overflow-x: auto;
        margin-top: 12px;
        border: 1px solid #e5e7eb;
        border-radius: 9px;
    }
    table.preview {
        width: 100%;
        border-collapse: collapse;
        font-size: 12px;
        min-width: 680px;
    }
    table.preview th, table.preview td {
        border-bottom: 1px solid #f1f5f9;
        padding: 8px 10px;
        text-align: right;
        vertical-align: top;
    }
    table.preview th {
        background: var(--brand-soft);
        color: #7c2d12;
        font-weight: 700;
        position: sticky;
        top: 0;
    }
    table.preview tr:hover td { background: #fffaf5; }
    .phone-chip {
        display: inline-block;
        direction: ltr;
        background: #ecfdf5;
        color: #065f46;
        border-radius: 5px;
        padding: 1px 6px;
        font-weight: 600;
        margin: 1px;
    }
    .land-chip {
        display: inline-block;
        direction: ltr;
        background: #eff6ff;
        color: #1e40af;
        border-radius: 5px;
        padding: 1px 6px;
        margin: 1px;
    }
    .counter-tag {
        font-size: 11.5px;
        font-weight: 600;
        color: #9a3412;
        background: var(--brand-soft);
        padding: 2px 9px;
        border-radius: 999px;
    }
</style>
</head>
<body>

<header>
    <h1>🗺️ استخراج‌کنندهٔ هوشمند اطلاعات کسب‌وکارها از گوگل مپ</h1>
    <p>جستجوی خودکار چندکلمه‌ای و چندشهری • پیش‌پالایش فوق‌سریع تکراری‌ها • خروجی اکسل و CSV در لحظه</p>
    <div class="version-pill">نسخهٔ تک‌فایلی مستقل ۳.۰ — بدون نیاز به هیچ فایل جانبی</div>
</header>

<div class="container">

    <!-- ۱. تنظیمات جستجو -->
    <div class="card">
        <h2>
            <span>۱. تنظیمات جستجو و لیست لوکیشن‌ها</span>
            <span id="fileNameBadge" class="counter-tag" style="display:none;"></span>
        </h2>

        <div class="warn-box">
            <strong>💡 نکتهٔ مهم دربارهٔ لوکیشن‌ها:</strong>
            گوگل مپ مرز رسمی «بخش» یا «دهستان» را ندارد و آن‌ها را به نزدیک‌ترین شهر/روستای هم‌نام نگاشت می‌کند.
            برنامه همیشه عبارت را به فرمت بهینهٔ <b>«لوکیشن + (کلمه۱ OR کلمه۲ ...)»</b> می‌سازد تا نتایج دقیقاً از همان منطقه برگردند و هرگز به شهر اشتباه (مثلاً تهران) نپرند.
        </div>

        <label>
            کلمات کلیدی جستجو (هر خط = یک عبارت کلیدی، حداکثر ۳۰ مورد)
            <span id="kwCountTag" class="counter-tag" style="margin-right:6px;">۰ / ۳۰</span>
        </label>
        <textarea id="keywords" style="min-height:105px;" placeholder="لوازم یدکی&#10;قطعات خودرو&#10;فروشگاه لوازم یدکی&#10;بازار لوازم یدکی"></textarea>
        <div class="hint">
            تمام خطوط با عملگر استاندارد <b>OR</b> ترکیب می‌شوند تا هر کسب‌وکاری که با هرکدام از این عناوین ثبت شده پیدا شود.
        </div>
        <div class="btn-row" style="margin-top:8px;">
            <button type="button" class="btn-secondary btn-sm" id="saveKeywordsBtn">💾 ذخیرهٔ کلمات کلیدی</button>
            <button type="button" class="btn-danger btn-sm" id="clearKeywordsBtn">پاک‌کردن کلمات کلیدی</button>
        </div>
        <div id="keywordsMsg" class="msg-line"></div>

        <label style="margin-top:16px;">
            لیست لوکیشن‌ها (هر خط = یک لوکیشن کامل)
            <span id="locCountTag" class="counter-tag" style="margin-right:6px;">۰ لوکیشن</span>
        </label>
        <textarea id="locations" placeholder="استان یزد شهرستان یزد بخش مرکزی دهستان فجر&#10;استان یزد شهرستان اردکان بخش مرکزی روستای عقدا&#10;استان تهران شهرستان تهران منطقه ۵"></textarea>
        <div class="hint">
            هنگام ذخیره، کاراکتر «/» خودکار با فاصله جایگزین شده و خطوط خالی یا تکراری حذف می‌شوند.
        </div>
        <div class="btn-row" style="margin-top:8px;">
            <button type="button" class="btn-secondary btn-sm" id="saveLocationsBtn">💾 ذخیره و بهینه‌سازی لوکیشن‌ها</button>
            <button type="button" class="btn-danger btn-sm" id="clearLocationsBtn">پاک‌کردن همهٔ لوکیشن‌ها</button>
        </div>
        <div id="locationsMsg" class="msg-line"></div>

        <div class="upload-box">
            <label style="margin-top:0;">📂 بارگذاری لیست لوکیشن‌ها از فایل اکسل (.xlsx) یا متنی (.csv / .txt)</label>
            <div class="row" style="align-items:center;">
                <div style="flex:2; min-width:220px;">
                    <input type="file" id="xlsxFile" accept=".xlsx,.xlsm,.csv,.txt">
                </div>
                <div style="flex:1; min-width:170px;">
                    <button type="button" class="btn-secondary btn-sm" id="uploadAppendBtn" style="width:100%;">➕ افزودن به انتهای لیست</button>
                </div>
                <div style="flex:1; min-width:170px;">
                    <button type="button" class="btn-secondary btn-sm" id="uploadReplaceBtn" style="width:100%;">🔄 جایگزینی کل لیست</button>
                </div>
            </div>
            <div class="hint">
                در فایل‌های چندستونی (مثلاً ستون استان، شهرستان، بخش، دهستان جدا)، سلول‌های هر ردیف خودکار با فاصله ترکیب می‌شوند و فایل خروجی هم همنام فایل ورودی (<code>نام‌فایل_نتایج.xlsx</code>) دانلود می‌شود.
            </div>
            <div id="uploadMsg" class="msg-line"></div>
        </div>

        <div class="row" style="margin-top:14px;">
            <div>
                <label>حداکثر نتیجه در هر لوکیشن</label>
                <input type="number" id="maxResults" value="150" min="1" max="600">
            </div>
            <div>
                <label>مکث بین اسکرول‌ها (ثانیه)</label>
                <input type="number" id="scrollPause" value="1.3" step="0.1" min="0.4">
            </div>
            <div>
                <label>مکث بین هر کسب‌وکار (ثانیه)</label>
                <input type="number" id="itemDelay" value="1.2" step="0.1" min="0.3">
            </div>
            <div>
                <label>مکث بین لوکیشن‌ها (ثانیه)</label>
                <input type="number" id="locDelay" value="4" step="0.5" min="0.5">
            </div>
        </div>

        <div class="check-grid">
            <label class="check-item">
                <input type="checkbox" id="skipExistingChk" checked>
                <span><b>رد کردن فوق‌سریع تکراری‌ها:</b> مکان‌هایی که قبلاً ثبت شده‌اند دوباره باز نمی‌شوند (۵ برابر سرعت بیشتر در شهرهای هم‌پوشان و Resume فوری).</span>
            </label>
            <label class="check-item">
                <input type="checkbox" id="blockMediaChk" checked>
                <span><b>حالت کم‌مصرف و سریع (توصیه‌شده):</b> عدم دانلود تصاویر سنگین نقشه برای کاهش ۶۰٪ مصرف رم و سرعت بالاتر.</span>
            </label>
            <label class="check-item">
                <input type="checkbox" id="headlessChk" checked>
                <span><b>اجرای مخفی مرورگر (Headless):</b> بدون باز شدن پنجرهٔ کروم (روی سرور/داکر حتماً فعال بماند).</span>
            </label>
        </div>

        <div class="btn-row" style="margin-top:18px;">
            <button class="btn-primary" id="startFreshBtn">🚀 شروع از ابتدا</button>
            <button class="btn-success" id="resumeBtn">▶️ ادامهٔ کار قبلی</button>
            <button class="btn-danger" id="stopBtn">⏹️ توقف اجرا</button>
        </div>
    </div>

    <!-- ۲. وضعیت اجرا -->
    <div class="card">
        <h2>
            <span>۲. وضعیت لحظه‌ای اجرا و لاگ زنده</span>
            <span id="healthBadge" class="health-pill">در حال اتصال...</span>
        </h2>

        <div class="status-header">
            <div>وضعیت فعلی: <span id="statusBadge" class="badge idle">غیرفعال</span></div>
            <label style="margin:0; font-weight:normal; font-size:12px; display:flex; align-items:center; gap:6px; cursor:pointer;">
                <input type="checkbox" id="autoScrollChk" checked> اسکرول خودکار لاگ
            </label>
        </div>

        <div class="status-box">
            <div class="stat"><div class="num" id="statLoc">۰ / ۰</div><div class="lbl">لوکیشن تکمیل‌شده</div></div>
            <div class="stat"><div class="num" id="statResults">۰</div><div class="lbl">کل کسب‌وکار یکتا</div></div>
            <div class="stat"><div class="num" id="statMobiles">۰</div><div class="lbl">دارای شماره موبایل (09)</div></div>
            <div class="stat"><div class="num" id="statPhones">۰</div><div class="lbl">دارای تلفن تماس</div></div>
            <div class="stat"><div class="num" id="statPct">۰٪</div><div class="lbl">پیشرفت کل لیست</div></div>
        </div>

        <div id="activeLocBox" class="current-loc-banner">لوکیشن فعال: —</div>

        <div class="progress-bar-outer"><div class="progress-bar-inner" id="progressBar"></div></div>

        <div id="errorBox" style="display:none;" class="warn-box error-box"></div>

        <div id="logBox">هنوز کاری اجرا نشده است...</div>
    </div>

    <!-- ۳. خروجی و پیش‌نمایش زنده -->
    <div class="card">
        <h2>
            <span>۳. دریافت خروجی اکسل / CSV و پیش‌نمایش زنده</span>
            <button type="button" class="btn-secondary btn-sm" id="refreshPreviewBtn">🔄 بروزرسانی جدول</button>
        </h2>

        <div class="row">
            <div>
                <label>قالب ستون‌های خروجی</label>
                <select id="exportLayout">
                    <option value="full">کامل و پیشرفته (۱۱ ستون: نام، دسته‌بندی، موبایل، ثابت، آدرس، وب‌سایت، امتیاز، مختصات، لینک)</option>
                    <option value="classic">کلاسیک مشابه نسخهٔ اصلی گیت‌هاب (۷ ستون اصلی)</option>
                </select>
            </div>
            <div>
                <label>فیلتر ردیف‌ها هنگام دانلود</label>
                <select id="exportFilter">
                    <option value="all">همهٔ کسب‌وکارهای استخراج‌شده</option>
                    <option value="with_phone">فقط موارد دارای شماره تماس (موبایل یا ثابت)</option>
                    <option value="with_mobile">فقط موارد دارای شماره موبایل (۰۹)</option>
                </select>
            </div>
        </div>

        <div class="btn-row">
            <button class="btn-primary" id="exportXlsxBtn">📥 دانلود فایل اکسل (.xlsx)</button>
            <button class="btn-secondary" id="exportCsvBtn">📄 دانلود فایل CSV (مناسب اکسل/CRM)</button>
        </div>
        <div class="hint">
            حتی زمانی که استخراج هنوز در حال اجراست، می‌توانید هر لحظه خروجی کامل نتایج ذخیره‌شده تا همین ثانیه را دانلود کنید.
        </div>

        <div style="margin-top:18px;">
            <input type="text" id="previewSearch" placeholder="🔍 جستجوی فوری در بین نتایج استخراج‌شده (نام، شهر، آدرس، شماره موبایل)...">
        </div>

        <div class="table-wrap">
            <table class="preview">
                <thead>
                    <tr>
                        <th>#</th>
                        <th>نام کسب‌وکار</th>
                        <th>دسته‌بندی</th>
                        <th>موبایل / تلفن</th>
                        <th>لوکیشن / آدرس</th>
                        <th>امتیاز</th>
                    </tr>
                </thead>
                <tbody id="previewBody">
                    <tr><td colspan="6" style="text-align:center;color:#888;padding:18px;">هنوز نتیجه‌ای ثبت نشده است.</td></tr>
                </tbody>
            </table>
        </div>
    </div>

    <!-- ۴. ریست کامل -->
    <div class="card" style="border-color:#fecaca;">
        <h2 style="color:#991b1b; border-bottom-color:#fecaca;">⚠️ ریست کامل برنامه</h2>
        <p style="font-size:13px; color:#4b5563; margin-top:0;">
            با زدن این دکمه (پس از دو مرحله تأیید)، تمام کلمات کلیدی، لیست لوکیشن‌ها، وضعیت پیشرفت و کلیهٔ نتایج استخراج‌شده از دیتابیس پاک می‌شوند. این عملیات <b>غیرقابل بازگشت</b> است.
        </p>
        <button type="button" class="btn-danger btn-sm" id="resetAllBtn">🗑️ ریست کامل و پاک‌سازی همهٔ اطلاعات</button>
        <div id="resetMsg" class="msg-line"></div>
    </div>

</div>

<script>
const $ = (id) => document.getElementById(id);

let logsSeen = 0;
let pollTimer = null;
let pollInFlight = false;
let lastPreviewCount = -1;
let MAX_KEYWORDS = 30;

function persianDigits(n) {
    return String(n ?? 0).replace(/\d/g, d => "۰۱۲۳۴۵۶۷۸۹"[d]);
}

function escapeHtml(str) {
    return String(str || "")
        .replace(/&/g, "&amp;")
        .replace(/</g, "&lt;")
        .replace(/>/g, "&gt;")
        .replace(/"/g, "&quot;");
}

function getKeywordsFromBox() {
    return $("keywords").value.split("\n").map(s => s.trim()).filter(Boolean);
}

function getLocationsFromBox() {
    return $("locations").value.split("\n").map(s => s.trim()).filter(Boolean);
}

function updateInputCounters() {
    const kwLen = getKeywordsFromBox().length;
    const locLen = getLocationsFromBox().length;
    $("kwCountTag").textContent = `${persianDigits(kwLen)} / ${persianDigits(MAX_KEYWORDS)}`;
    $("kwCountTag").style.color = kwLen > MAX_KEYWORDS ? "#b91c1c" : "#9a3412";
    $("locCountTag").textContent = `${persianDigits(locLen)} لوکیشن`;
}

function updateFileBadge(basename) {
    const el = $("fileNameBadge");
    if (basename) {
        el.style.display = "inline-block";
        el.textContent = `خروجی: ${basename}_نتایج.xlsx`;
    } else {
        el.style.display = "none";
        el.textContent = "";
    }
}

async function startJob(resume) {
    const keywords = getKeywordsFromBox();
    const locations = getLocationsFromBox();
    const max_results = parseInt($("maxResults").value || "150", 10);
    const scroll_pause = parseFloat($("scrollPause").value || "1.3");
    const item_delay = parseFloat($("itemDelay").value || "1.2");
    const between_location_delay = parseFloat($("locDelay").value || "4");
    const skip_existing = $("skipExistingChk").checked;
    const block_media = $("blockMediaChk").checked;
    const headless = $("headlessChk").checked;

    if (!keywords.length) { alert("لطفاً حداقل یک کلمهٔ کلیدی وارد کنید."); return; }
    if (keywords.length > MAX_KEYWORDS) {
        alert(`حداکثر ${persianDigits(MAX_KEYWORDS)} کلمهٔ کلیدی مجاز است (شما ${persianDigits(keywords.length)} مورد وارد کرده‌اید).`);
        return;
    }
    if (!locations.length) { alert("لطفاً حداقل یک لوکیشن وارد کنید."); return; }

    if (!resume && lastPreviewCount > 0) {
        const confirmWipe = confirm(
            `هم‌اکنون ${persianDigits(lastPreviewCount)} نتیجه از قبل در دیتابیس ذخیره شده است.\n` +
            `با زدن «شروع از ابتدا»، نتایج قبلی پاک شده و کار از صفر شروع می‌شود.\n` +
            `(اگر می‌خواهید نتایج قبلی حفظ شود، Cancel را بزنید و «ادامهٔ کار قبلی» را انتخاب کنید.)\n\n` +
            `آیا از شروع مجدد از صفر مطمئن هستید؟`
        );
        if (!confirmWipe) return;
    }

    if (locations.length > 300) {
        const ok = confirm(
            `تعداد لوکیشن‌های واردشده ${persianDigits(locations.length)} مورد است و ممکن است ساعت‌ها یا روزها زمان ببرد.\nآیا ادامه می‌دهید؟`
        );
        if (!ok) return;
    }

    logsSeen = 0;
    $("logBox").textContent = "";

    const res = await fetch("/api/start", {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify({
            keywords, locations, max_results, resume,
            scroll_pause, item_delay, between_location_delay,
            skip_existing, block_media, headless
        })
    });
    const data = await res.json();
    if (!data.ok) {
        alert("خطا: " + data.error);
        return;
    }
    startPolling();
}

async function stopJob() {
    const res = await fetch("/api/stop", { method: "POST" });
    const data = await res.json();
    if (!data.ok) alert("خطا: " + data.error);
    pollStatus();
}

function exportResults(fmt) {
    const layout = $("exportLayout").value;
    const filter = $("exportFilter").value;
    window.location.href = `/api/export?format=${encodeURIComponent(fmt)}&layout=${encodeURIComponent(layout)}&filter=${encodeURIComponent(filter)}`;
}

async function loadPreview() {
    try {
        const q = $("previewSearch").value.trim();
        const res = await fetch(`/api/results_preview?q=${encodeURIComponent(q)}`);
        const data = await res.json();
        const tbody = $("previewBody");
        if (!data.sample || !data.sample.length) {
            tbody.innerHTML = `<tr><td colspan="6" style="text-align:center;color:#888;padding:18px;">${
                q ? "موردی مطابق جستجوی شما یافت نشد." : "هنوز نتیجه‌ای ثبت نشده است."
            }</td></tr>`;
            return;
        }
        tbody.innerHTML = data.sample.map(r => {
            const mobs = (r.mobiles || []).map(m => `<span class="phone-chip">📱 ${escapeHtml(m)}</span>`).join(" ");
            const lands = (r.landlines || []).map(l => `<span class="land-chip">☎️ ${escapeHtml(l)}</span>`).join(" ");
            const phones = (mobs || lands) ? `${mobs} ${lands}` : `<span style="color:#9ca3af;">—</span>`;
            const nameCell = r.url
                ? `<a href="${escapeHtml(r.url)}" target="_blank" rel="noopener" style="color:#9a3412;font-weight:700;text-decoration:none;">${escapeHtml(r.name)} ↗</a>`
                : `<b>${escapeHtml(r.name)}</b>`;
            return `<tr>
                <td>${persianDigits(r.id)}</td>
                <td>${nameCell}</td>
                <td>${escapeHtml(r.category || "—")}</td>
                <td>${phones}</td>
                <td><div style="font-weight:600;color:#374151;">${escapeHtml(r.location || "")}</div><div style="color:#6b7280;font-size:11px;">${escapeHtml(r.address || "")}</div></td>
                <td>${r.rating ? "⭐ " + persianDigits(r.rating) : "—"}</td>
            </tr>`;
        }).join("");
    } catch (_e) {}
}

function startPolling() {
    if (pollTimer) clearInterval(pollTimer);
    pollTimer = setInterval(pollStatus, 1500);
    pollStatus();
}

async function pollStatus() {
    if (pollInFlight) return;
    pollInFlight = true;
    let data;
    try {
        const res = await fetch(`/api/status?logs_since=${logsSeen}`);
        data = await res.json();
    } catch (_e) {
        pollInFlight = false;
        return;
    }
    pollInFlight = false;

    const badge = $("statusBadge");
    if (data.error) {
        badge.textContent = "خطا / متوقف‌شده";
        badge.className = "badge error";
        $("errorBox").style.display = "block";
        $("errorBox").textContent = data.error;
    } else if (data.running && data.stop_requested) {
        badge.textContent = "در حال توقف امن...";
        badge.className = "badge stopping";
        $("errorBox").style.display = "none";
    } else if (data.running) {
        badge.textContent = "در حال استخراج...";
        badge.className = "badge running";
        $("errorBox").style.display = "none";
    } else {
        badge.textContent = "آماده به کار";
        badge.className = "badge idle";
        $("errorBox").style.display = "none";
    }

    $("startFreshBtn").disabled = !!data.running;
    $("resumeBtn").disabled = !!data.running;
    $("stopBtn").disabled = !data.running;

    const done = Math.max(0, (data.current_index ?? -1) + 1);
    const total = data.total_locations || 0;
    const pct = total > 0 ? Math.min(100, Math.round((done / total) * 100)) : 0;

    $("statLoc").textContent = `${persianDigits(done)} / ${persianDigits(total)}`;
    $("statResults").textContent = persianDigits(data.results_count || 0);
    $("statMobiles").textContent = persianDigits(data.mobile_count || 0);
    $("statPhones").textContent = persianDigits(data.phone_count || 0);
    $("statPct").textContent = `${persianDigits(pct)}٪`;
    $("progressBar").style.width = pct + "%";

    if (data.active_location && data.running) {
        $("activeLocBox").textContent = `📍 در حال بررسی لوکیشن: ${data.active_location}`;
    } else if (total > 0 && done >= total) {
        $("activeLocBox").textContent = `✅ پردازش تمام ${persianDigits(total)} لوکیشن به پایان رسید.`;
    } else {
        $("activeLocBox").textContent = "لوکیشن فعال: —";
    }

    if (data.health) {
        const h = data.health;
        const ramStr = h.rss_mb ? `RAM: ${h.rss_mb}MB` : "RAM: OK";
        $("healthBadge").textContent = `${ramStr} | Threads: ${h.threads} | Chrome: ${h.chrome_procs}`;
    }

    if (data.logs_total < logsSeen) {
        $("logBox").textContent = "";
        logsSeen = 0;
    }
    if (data.logs && data.logs.length) {
        const box = $("logBox");
        box.textContent += data.logs.join("\n") + "\n";
        if ($("autoScrollChk").checked) {
            box.scrollTop = box.scrollHeight;
        }
    }
    logsSeen = data.logs_total;

    if (data.results_count !== lastPreviewCount) {
        lastPreviewCount = data.results_count;
        loadPreview();
    }

    if (!data.running && pollTimer) {
        clearInterval(pollTimer);
        pollTimer = setInterval(pollStatus, 4500);
    }
}

async function loadSettings() {
    try {
        const res = await fetch("/api/settings");
        const data = await res.json();
        if (!data.ok) return;
        MAX_KEYWORDS = data.max_keywords || 30;
        if (data.keywords && data.keywords.length) $("keywords").value = data.keywords.join("\n");
        if (data.locations && data.locations.length) $("locations").value = data.locations.join("\n");
        if (data.max_results) $("maxResults").value = data.max_results;
        if (data.scroll_pause) $("scrollPause").value = data.scroll_pause;
        if (data.item_delay) $("itemDelay").value = data.item_delay;
        if (data.between_location_delay) $("locDelay").value = data.between_location_delay;
        if (typeof data.skip_existing === "boolean") $("skipExistingChk").checked = data.skip_existing;
        if (typeof data.block_media === "boolean") $("blockMediaChk").checked = data.block_media;
        if (typeof data.headless === "boolean") $("headlessChk").checked = data.headless;
        updateFileBadge(data.location_file_basename || "");
        updateInputCounters();
    } catch (_e) {}
}

async function saveKeywords() {
    const keywords = getKeywordsFromBox();
    const msg = $("keywordsMsg");
    if (keywords.length > MAX_KEYWORDS) {
        msg.textContent = `حداکثر ${persianDigits(MAX_KEYWORDS)} کلمهٔ کلیدی مجاز است.`;
        return;
    }
    const res = await fetch("/api/settings/keywords", {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify({ keywords })
    });
    const data = await res.json();
    msg.textContent = data.ok
        ? `✅ ${persianDigits(keywords.length)} کلمهٔ کلیدی ذخیره شد.`
        : "خطا: " + data.error;
    updateInputCounters();
}

async function clearKeywords() {
    if (!confirm("همهٔ کلمات کلیدی پاک شوند؟")) return;
    const res = await fetch("/api/settings/clear", {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify({ what: "keywords" })
    });
    const data = await res.json();
    if (data.ok) {
        $("keywords").value = "";
        $("keywordsMsg").textContent = "همهٔ کلمات کلیدی پاک شدند.";
        updateInputCounters();
    }
}

async function saveLocations() {
    const locations = getLocationsFromBox();
    const msg = $("locationsMsg");
    const res = await fetch("/api/settings/locations", {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify({ locations })
    });
    const data = await res.json();
    if (!data.ok) { msg.textContent = "خطا: " + data.error; return; }
    $("locations").value = data.locations.join("\n");
    msg.textContent = `✅ ${persianDigits(data.count)} لوکیشن یکتا بهینه‌سازی و ذخیره شد.`;
    updateInputCounters();
}

async function clearLocations() {
    if (!confirm("همهٔ لوکیشن‌ها پاک شوند؟")) return;
    const res = await fetch("/api/settings/clear", {
        method: "POST",
        headers: { "Content-Type": "application/json" },
        body: JSON.stringify({ what: "locations" })
    });
    const data = await res.json();
    if (data.ok) {
        $("locations").value = "";
        $("locationsMsg").textContent = "همهٔ لوکیشن‌ها پاک شدند.";
        updateFileBadge("");
        updateInputCounters();
    }
}

async function uploadLocationsFile(replace) {
    const fileInput = $("xlsxFile");
    const msg = $("uploadMsg");
    if (!fileInput.files || !fileInput.files.length) {
        msg.textContent = "ابتدا یک فایل اکسل یا متنی انتخاب کنید.";
        return;
    }
    const formData = new FormData();
    formData.append("file", fileInput.files[0]);
    formData.append("mode", replace ? "replace" : "append");

    msg.textContent = "در حال خواندن و بهینه‌سازی فایل...";
    try {
        const res = await fetch("/api/upload_locations", { method: "POST", body: formData });
        const data = await res.json();
        if (!data.ok) {
            msg.textContent = "خطا: " + data.error;
            return;
        }
        $("locations").value = data.all_locations.join("\n");
        updateFileBadge(data.location_file_basename || "");
        updateInputCounters();
        const fname = data.location_file_basename;
        msg.textContent = fname
            ? `✅ ${persianDigits(data.added_count)} ردیف خوانده شد (مجموع یکتا: ${persianDigits(data.total_count)}) و خودکار ذخیره شد. خروجی با نام «${fname}_نتایج.xlsx» دانلود می‌شود.`
            : `✅ ${persianDigits(data.added_count)} لوکیشن خوانده و ذخیره شد.`;
    } catch (e) {
        msg.textContent = "خطا در ارسال فایل: " + e;
    }
}

async function resetAll() {
    if (!confirm("⚠️ هشدار: تمام کلمات کلیدی، لیست لوکیشن‌ها و کلیهٔ نتایج استخراج‌شده برای همیشه پاک می‌شوند. ادامه می‌دهید؟")) return;
    if (!confirm("تأیید نهایی: آیا مطمئن هستید که کل دیتابیس برنامه ریست شود؟")) return;

    const msg = $("resetMsg");
    msg.textContent = "در حال پاک‌سازی کامل دیتابیس...";
    const res = await fetch("/api/reset", { method: "POST" });
    const data = await res.json();
    if (!data.ok) { msg.textContent = "خطا: " + data.error; return; }
    msg.textContent = "برنامه به‌طور کامل ریست شد. در حال بارگذاری مجدد...";
    setTimeout(() => location.reload(), 900);
}

let searchDebounce = null;
$("previewSearch").addEventListener("input", () => {
    clearTimeout(searchDebounce);
    searchDebounce = setTimeout(loadPreview, 250);
});

$("keywords").addEventListener("input", updateInputCounters);
$("locations").addEventListener("input", updateInputCounters);

$("uploadAppendBtn").addEventListener("click", () => uploadLocationsFile(false));
$("uploadReplaceBtn").addEventListener("click", () => uploadLocationsFile(true));
$("saveKeywordsBtn").addEventListener("click", saveKeywords);
$("clearKeywordsBtn").addEventListener("click", clearKeywords);
$("saveLocationsBtn").addEventListener("click", saveLocations);
$("clearLocationsBtn").addEventListener("click", clearLocations);
$("refreshPreviewBtn").addEventListener("click", loadPreview);
$("resetAllBtn").addEventListener("click", resetAll);

$("startFreshBtn").addEventListener("click", () => startJob(false));
$("resumeBtn").addEventListener("click", () => startJob(true));
$("stopBtn").addEventListener("click", stopJob);
$("exportXlsxBtn").addEventListener("click", () => exportResults("xlsx"));
$("exportCsvBtn").addEventListener("click", () => exportResults("csv"));

loadSettings();
loadPreview();
startPolling();
</script>

</body>
</html>
"""


# ==============================================================================
# اپلیکیشن وب Flask و مدیریت Thread اجرای کار پس‌زمینه
# ==============================================================================

storage = StorageManager(DB_PATH)
app = Flask(__name__)

job_lock = threading.Lock()
job_state = {
    "running": False,
    "stop_requested": False,
    "keyword": "",
    "locations": [],
    "current_index": -1,
    "active_location": "",
    "total_locations": 0,
    "logs": [],
    "started_at": None,
    "finished_at": None,
    "error": None,
}
scraper_instance: GoogleMapsScraper | None = None


def _restore_initial_job_state() -> None:
    saved = storage.get_job_state()
    if saved:
        with job_lock:
            job_state.update(saved)
            job_state["running"] = False
            job_state["stop_requested"] = False
            job_state["active_location"] = ""
            job_state["logs"] = []


_restore_initial_job_state()


def log(msg: str) -> None:
    ts = time.strftime("%H:%M:%S")
    line = f"[{ts}] {msg}"
    print(line, flush=True)
    with job_lock:
        job_state["logs"].append(line)
        if len(job_state["logs"]) > 2500:
            job_state["logs"] = job_state["logs"][-2500:]


def run_job(
    keyword: str,
    locations: list[str],
    max_results: int,
    resume: bool,
    scroll_pause: float,
    item_delay: float,
    between_location_delay: float,
    skip_existing: bool,
    block_media: bool,
    headless: bool,
) -> None:
    global scraper_instance

    with job_lock:
        job_state["running"] = True
        job_state["stop_requested"] = False
        job_state["keyword"] = keyword
        job_state["locations"] = locations
        job_state["total_locations"] = len(locations)
        job_state["started_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
        job_state["finished_at"] = None
        job_state["error"] = None
        job_state["logs"] = []

        if resume:
            start_index = max(0, job_state["current_index"] + 1)
        else:
            job_state["current_index"] = -1
            start_index = 0

    if not resume:
        storage.clear_results_only()

    storage.save_job_state(job_state)

    log(
        f"شروع عملیات استخراج | عبارت جستجو: «{keyword}» | تعداد لوکیشن‌ها: {len(locations)} | "
        f"حالت: {'ادامهٔ کار قبلی' if resume else 'شروع تازه از ابتدا'}"
    )
    if start_index >= len(locations) and resume:
        log("تمامی لوکیشن‌های این لیست قبلاً به پایان رسیده‌اند! برای اجرای مجدد از «شروع از ابتدا» استفاده کنید.")

    log(health_log_line(storage))

    scraper = GoogleMapsScraper(
        headless=headless,
        block_media=block_media,
        log_fn=log,
        seen_checker=storage.has_seen,
        storage_ref=storage,
    )
    scraper_instance = scraper

    try:
        if start_index < len(locations):
            scraper.start()

        for idx in range(start_index, len(locations)):
            with job_lock:
                if job_state["stop_requested"]:
                    log("عملیات به درخواست کاربر متوقف شد.")
                    break
                job_state["active_location"] = f"({idx + 1}/{len(locations)}) {locations[idx]}"

            location = locations[idx]
            log(f"═══ لوکیشن {idx + 1} از {len(locations)}: {location} ═══")

            try:
                saved_count = scraper.search_location(
                    keyword=keyword,
                    location=location,
                    max_results=max_results,
                    scroll_pause=scroll_pause,
                    item_delay=item_delay,
                    skip_existing=skip_existing,
                    on_item_extracted=storage.append_result,
                )
                log(
                    f"پایان لوکیشن «{location}»: {saved_count} کسب‌وکار جدید ذخیره شد "
                    f"(مجموع کل تا این لحظه: {storage.total_count})."
                )
            except RuntimeError as exc:
                log(f"⛔ توقف ایمن به‌دلیل محدودیت/کپچا: {exc}")
                with job_lock:
                    job_state["error"] = str(exc)
                break
            except Exception as exc:
                log(f"⚠️ خطا در پردازش لوکیشن «{location}»: {exc}")
                log(traceback.format_exc(limit=3))

            with job_lock:
                stop_now = job_state["stop_requested"]

            # فقط اگر لوکیشن بدون توقف دستی کامل شده باشد ایندکس جلو می‌رود تا
            # در «ادامهٔ کار قبلی» لوکیشن نیمه‌کاره از قلم نیفتد (آیتم‌های استخراج‌شدهٔ
            # آن لوکیشن به‌کمک پیش‌پالایش هوشمند در ۰ میلی‌ثانیه رد می‌شوند).
            if not stop_now:
                with job_lock:
                    job_state["current_index"] = idx
                storage.save_job_state(job_state)
            else:
                log("توقف کامل انجام شد. پیشرفت شما ذخیره شده و هر زمان با «ادامهٔ کار قبلی» قابل ادامه است.")
                break

            log(health_log_line(storage))

            if idx < len(locations) - 1:
                time.sleep(between_location_delay)

    except Exception as exc:
        log(f"❌ خطای غیرمنتظره در اجرای جاب: {exc}")
        log(traceback.format_exc(limit=5))
        with job_lock:
            job_state["error"] = str(exc)
    finally:
        try:
            scraper.stop()
        except Exception:
            pass
        scraper_instance = None
        with job_lock:
            job_state["running"] = False
            job_state["active_location"] = ""
            job_state["finished_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
        storage.save_job_state(job_state)
        log("🏁 پایان اجرای فعلی.")


# ------------------------------------------------------------------------------
# مسیرهای API وب‌سرور
# ------------------------------------------------------------------------------

@app.route("/")
def index():
    return Response(HTML_TEMPLATE, mimetype="text/html; charset=utf-8")


@app.route("/api/health", methods=["GET"])
def api_health():
    return jsonify({"ok": True, **health_snapshot(storage)})


@app.route("/api/settings", methods=["GET"])
def api_get_settings():
    settings = storage.get_settings()
    return jsonify({"ok": True, **settings, "max_keywords": MAX_KEYWORDS})


@app.route("/api/settings/keywords", methods=["POST"])
def api_save_keywords():
    data = request.get_json(force=True) or {}
    raw_keywords = data.get("keywords") or []
    keywords = [str(k).strip() for k in raw_keywords if str(k).strip()]
    if len(keywords) > MAX_KEYWORDS:
        return jsonify({
            "ok": False,
            "error": f"حداکثر {MAX_KEYWORDS} کلمهٔ کلیدی مجاز است (شما {len(keywords)} مورد وارد کرده‌اید).",
        }), 400

    settings = storage.get_settings()
    settings["keywords"] = keywords
    storage.save_settings(settings)
    return jsonify({"ok": True, "keywords": keywords})


@app.route("/api/settings/locations", methods=["POST"])
def api_save_locations():
    data = request.get_json(force=True) or {}
    raw_locations = data.get("locations") or []
    locations = [str(l).strip() for l in raw_locations if str(l).strip()]
    optimized = optimize_locations(locations)

    settings = storage.get_settings()
    settings["locations"] = optimized
    storage.save_settings(settings)
    return jsonify({"ok": True, "locations": optimized, "count": len(optimized)})


@app.route("/api/settings/clear", methods=["POST"])
def api_clear_settings():
    data = request.get_json(force=True) or {}
    what = data.get("what", "all")

    settings = storage.get_settings()
    if what in ("keywords", "all"):
        settings["keywords"] = []
    if what in ("locations", "all"):
        settings["locations"] = []
        settings["location_file_basename"] = ""
    storage.save_settings(settings)
    return jsonify({"ok": True, **settings})


@app.route("/api/upload_locations", methods=["POST"])
def api_upload_locations():
    """پشتیبانی از آپلود فایل اکسل (.xlsx / .xlsm) و فایل متنی/CSV (.csv / .txt)
    همراه با ذخیرهٔ خودکار روی دیسک و نام‌گذاری خودکار خروجی."""
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "هیچ فایلی ارسال نشد."}), 400

    uploaded = request.files["file"]
    mode = (request.form.get("mode") or "replace").strip().lower()
    raw_name = uploaded.filename or ""
    lower_name = raw_name.lower()

    if not lower_name.endswith((".xlsx", ".xlsm", ".csv", ".txt")):
        return jsonify({
            "ok": False,
            "error": "فقط فایل‌های اکسل (.xlsx / .xlsm) یا متنی (.csv / .txt) پشتیبانی می‌شوند.",
        }), 400

    extracted_rows: list[str] = []
    try:
        if lower_name.endswith((".xlsx", ".xlsm")):
            wb = load_workbook(uploaded, read_only=True, data_only=True)
            try:
                ws = wb.worksheets[0]
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c).strip() for c in row if c is not None and str(c).strip() != ""]
                    if cells:
                        extracted_rows.append(" ".join(cells))
            finally:
                wb.close()
        else:
            raw_bytes = uploaded.read()
            text = raw_bytes.decode("utf-8-sig", errors="replace")
            if lower_name.endswith(".csv"):
                for row in csv.reader(io.StringIO(text)):
                    cells = [str(c).strip() for c in row if c and str(c).strip()]
                    if cells:
                        extracted_rows.append(" ".join(cells))
            else:
                for line in text.splitlines():
                    if line.strip():
                        extracted_rows.append(line.strip())
    except Exception as exc:
        return jsonify({"ok": False, "error": f"خطا در خواندن فایل: {exc}"}), 400

    optimized_new = optimize_locations(extracted_rows)
    if not optimized_new:
        return jsonify({"ok": False, "error": "هیچ ردیف معتبری در فایل یافت نشد."}), 400

    basename = re.sub(r"\.(xlsx|xlsm|csv|txt)$", "", raw_name, flags=re.IGNORECASE).strip()
    basename = re.sub(r'[\\/:*?"<>|]', "_", basename)

    settings = storage.get_settings()
    if mode == "append":
        combined = optimize_locations(settings.get("locations", []) + optimized_new)
    else:
        combined = optimized_new

    settings["locations"] = combined
    settings["location_file_basename"] = basename
    storage.save_settings(settings)

    return jsonify({
        "ok": True,
        "added_count": len(optimized_new),
        "total_count": len(combined),
        "all_locations": combined,
        "location_file_basename": basename,
    })


@app.route("/api/start", methods=["POST"])
def api_start():
    data = request.get_json(force=True) or {}

    raw_keywords = data.get("keywords")
    if raw_keywords is not None:
        keywords = [str(k).strip() for k in raw_keywords if str(k).strip()]
        if len(keywords) > MAX_KEYWORDS:
            return jsonify({
                "ok": False,
                "error": f"حداکثر {MAX_KEYWORDS} کلمهٔ کلیدی مجاز است.",
            }), 400
        keyword = build_search_keyword(keywords)
    else:
        keywords = []
        keyword = (data.get("keyword") or "").strip()

    raw_locations = data.get("locations") or []
    locations = optimize_locations([str(loc).strip() for loc in raw_locations if str(loc).strip()])

    max_results = max(1, int(data.get("max_results") or 150))
    resume = bool(data.get("resume"))
    scroll_pause = max(0.3, float(data.get("scroll_pause") or 1.3))
    item_delay = max(0.2, float(data.get("item_delay") or 1.2))
    between_location_delay = max(0.5, float(data.get("between_location_delay") or 4.0))
    skip_existing = bool(data.get("skip_existing", True))
    block_media = bool(data.get("block_media", True))
    headless = bool(data.get("headless", True))

    if not keyword:
        return jsonify({"ok": False, "error": "حداقل یک کلمهٔ کلیدی جستجو وارد کنید."}), 400
    if not locations:
        return jsonify({"ok": False, "error": "حداقل یک لوکیشن معتبر وارد کنید."}), 400

    with job_lock:
        if job_state["running"]:
            return jsonify({"ok": False, "error": "یک عملیات هم‌اکنون در حال اجراست؛ ابتدا آن را متوقف کنید."}), 409

    settings = storage.get_settings()
    if keywords:
        settings["keywords"] = keywords
    settings["locations"] = locations
    settings["max_results"] = max_results
    settings["scroll_pause"] = scroll_pause
    settings["item_delay"] = item_delay
    settings["between_location_delay"] = between_location_delay
    settings["skip_existing"] = skip_existing
    settings["block_media"] = block_media
    settings["headless"] = headless
    storage.save_settings(settings)

    worker = threading.Thread(
        target=run_job,
        args=(
            keyword,
            locations,
            max_results,
            resume,
            scroll_pause,
            item_delay,
            between_location_delay,
            skip_existing,
            block_media,
            headless,
        ),
        daemon=True,
    )
    worker.start()
    return jsonify({"ok": True})


@app.route("/api/stop", methods=["POST"])
def api_stop():
    with job_lock:
        if not job_state["running"]:
            return jsonify({"ok": False, "error": "در حال حاضر عملیاتی در حال اجرا نیست."}), 400
        job_state["stop_requested"] = True
    if scraper_instance:
        scraper_instance.request_stop()
    log("🛑 درخواست توقف ثبت شد؛ مرورگر در حال توقف ایمن و ذخیرهٔ وضعیت است...")
    return jsonify({"ok": True})


@app.route("/api/status", methods=["GET"])
def api_status():
    since = max(0, int(request.args.get("logs_since", 0)))
    with job_lock:
        logs = job_state["logs"][since:]
        total_logs = len(job_state["logs"])
        snapshot = {
            "running": job_state["running"],
            "stop_requested": job_state["stop_requested"],
            "keyword": job_state["keyword"],
            "current_index": job_state["current_index"],
            "active_location": job_state["active_location"],
            "total_locations": job_state["total_locations"],
            "started_at": job_state["started_at"],
            "finished_at": job_state["finished_at"],
            "error": job_state["error"],
        }
    snapshot["results_count"] = storage.total_count
    snapshot["mobile_count"] = storage.mobile_count
    snapshot["phone_count"] = storage.phone_count
    snapshot["health"] = health_snapshot(storage)
    snapshot["logs"] = logs
    snapshot["logs_total"] = total_logs
    return jsonify(snapshot)


@app.route("/api/results_preview", methods=["GET"])
def api_results_preview():
    search_q = request.args.get("q", "")
    sample = storage.get_recent_results(limit=30, search_q=search_q)
    return jsonify({"count": storage.total_count, "sample": sample})


@app.route("/api/export", methods=["GET"])
def api_export():
    fmt = (request.args.get("format") or "xlsx").strip().lower()
    layout = (request.args.get("layout") or "full").strip().lower()
    filter_mode = (request.args.get("filter") or "all").strip().lower()

    records = storage.load_all_results(filter_mode=filter_mode)
    if not records:
        return jsonify({
            "ok": False,
            "error": "هیچ نتیجه‌ای مطابق فیلتر انتخابی برای خروجی یافت نشد.",
        }), 400

    settings = storage.get_settings()
    basename = settings.get("location_file_basename", "")
    suffix = ""
    if filter_mode == "with_mobile":
        suffix = "_موبایل"
    elif filter_mode == "with_phone":
        suffix = "_تلفن‌دار"

    stem = f"{basename}_نتایج{suffix}" if basename else f"نتایج-استخراج-گوگل-مپ{suffix}"

    if fmt == "csv":
        buf = build_csv_bytes(records, layout=layout)
        return send_file(
            buf,
            as_attachment=True,
            download_name=f"{stem}.csv",
            mimetype="text/csv; charset=utf-8",
        )

    buf = build_excel_bytes(records, layout=layout)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"{stem}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/reset", methods=["POST"])
def api_reset():
    with job_lock:
        if job_state["running"]:
            return jsonify({
                "ok": False,
                "error": "یک عملیات در حال اجراست. ابتدا آن را متوقف کرده و سپس ریست کنید.",
            }), 409

    storage.reset_everything()

    with job_lock:
        job_state.update({
            "running": False,
            "stop_requested": False,
            "keyword": "",
            "locations": [],
            "current_index": -1,
            "active_location": "",
            "total_locations": 0,
            "logs": [],
            "started_at": None,
            "finished_at": None,
            "error": None,
        })

    log("تمامی اطلاعات برنامه به‌طور کامل ریست شد.")
    return jsonify({"ok": True})


# ==============================================================================
# اجرای سرور پایدار Waitress (یا Werkzeug سخت‌گیرانه) + باز کردن مرورگر محلی
# ==============================================================================

def _is_container_or_headless_env() -> bool:
    if os.path.exists("/.dockerenv"):
        return True
    for env_key in ("RAILWAY_ENVIRONMENT", "DOCKER_CONTAINER", "KUBERNETES_SERVICE_HOST", "CODESPACES"):
        if os.environ.get(env_key):
            return True
    if sys.platform.startswith("linux") and not os.environ.get("DISPLAY") and not os.environ.get("WAYLAND_DISPLAY"):
        return True
    return False


def run_web_server(host: str, port: int, open_browser: bool = True) -> None:
    local_url = f"http://127.0.0.1:{port}"
    print("=" * 72, flush=True)
    print(f"🗺️  پنل استخراج گوگل مپ | نسخه: {APP_VERSION}", flush=True)
    print(f"🌐 آدرس دسترسی پنل وب : {local_url}", flush=True)
    print(f"💾 فایل پایگاه‌داده   : {DB_PATH}", flush=True)
    print(health_log_line(storage), flush=True)
    print("=" * 72, flush=True)

    if open_browser and not _is_container_or_headless_env():
        threading.Timer(1.2, lambda: webbrowser.open(local_url)).start()

    try:
        from waitress import serve

        print(f"[سرور] اجرای Waitress (threads=8, channel_timeout=120) روی {host}:{port}", flush=True)
        serve(app, host=host, port=port, threads=8, channel_timeout=120, ident="gmaps-scraper")
        return
    except ImportError:
        print("[سرور] Waitress یافت نشد؛ استفاده از Werkzeug مقاوم‌شده...", flush=True)

    from werkzeug.serving import ThreadedWSGIServer, WSGIRequestHandler

    class _HardenedHandler(WSGIRequestHandler):
        timeout = 60
        protocol_version = "HTTP/1.0"

    class _HardenedServer(ThreadedWSGIServer):
        daemon_threads = True

    _HardenedServer(host, port, app, handler=_HardenedHandler).serve_forever()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="پنل تک‌فایلی و هوشمند استخراج اطلاعات کسب‌وکارها از گوگل مپ"
    )
    parser.add_argument(
        "--host",
        default="0.0.0.0",
        help="آدرس شبکه برای گوش دادن وب‌سرور (پیش‌فرض: 0.0.0.0)",
    )
    parser.add_argument(
        "--port",
        type=int,
        default=int(os.environ.get("PORT", 5000)),
        help="پورت وب‌سرور (پیش‌فرض: 5000 یا مقدار متغیر محیطی PORT)",
    )
    parser.add_argument(
        "--no-browser",
        action="store_true",
        help="عدم باز کردن خودکار مرورگر هنگام اجرا روی ویندوز/مک",
    )
    parser.add_argument(
        "--install-browser",
        action="store_true",
        help="نصب خودکار مرورگر Chromium موردنیاز Playwright و خروج",
    )
    args = parser.parse_args()

    if args.install_browser:
        ok = install_playwright_chromium(print)
        sys.exit(0 if ok else 1)

    run_web_server(host=args.host, port=args.port, open_browser=not args.no_browser)


if __name__ == "__main__":
    main()
